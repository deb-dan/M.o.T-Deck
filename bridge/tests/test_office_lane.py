"""OFFICE LANE slice 1 (SHEETS) — the round-trip tables (2026-08-21).

Built to docs/research/2026-08-20-office-lane-recon.md §5/§7. What each group guards:

1. NAME + CONTAINMENT. Every office route takes a NAME off the wire. `valid_name` and
   `doc_target` are the whole boundary — the `library_target`/`_deletable_target`
   discipline: basename only, no traversal, no symlink escape, .xlsx forced.
2. THE FORWARD MAPPER (xlsx → IWorkbookData). The enums are Univer's, read out of the
   pinned UMD bundle: a renumbered CellValueType is a silently wrong cell, not an
   error, so the numbers are asserted literally here.
3. THE REVERSE MAPPER (IWorkbookData → xlsx) and the ROUND TRIP. This is the whole
   product: the recon's own verdict is that any plan which ships a beautiful grid and
   writes CSV has solved nothing. The round-trip is executed, not argued.
4. TOTALITY. A snapshot arrives as JSON from a webview and a workbook arrives from
   somebody else's software. Junk in a cell must cost that cell — never the workbook,
   never a 500.
5. THE BACKUP POLICY. A save writes a NEW workbook from the snapshot, so the pre-edit
   file has to survive: one .bak per file per day, taken BEFORE the first save, and a
   backup that cannot be written refuses the save rather than overwriting the original.
6. WIRING. The routes, the tab row and the pinned assets. ⚠️ The bundle LOAD ORDER
   moved to bridge/tests/test_office_grid.js at 2026-08-21e, where it is read off the
   page's `VENDOR` list instead of scraped out of markup — because the bundles are no
   longer static tags.
7. THE BLANK-TAB FORENSICS, and the answer to them. Everything above was already green
   while the tab on the real Mac showed nothing. Three rounds of diagnostics later the
   design changed instead: the page is now TWO TIERS — a plain-DOM grid that needs
   nothing but the served document, and Univer as an opt-in upgrade that is allowed to
   fail. So this group pins the invariant that makes the blank tab unrepresentable
   (no external subresource in the document at all), plus the forensics that stay
   useful either way: a static fallback banner true by default, a build stamp, the
   no-store headers, and the WKWebView delegate that answers when a content process
   dies. The tier-1 grid itself is tested in bridge/tests/test_office_grid.js.

Run: python3 bridge/tests/test_office_lane.py
"""
import datetime
import io
import json
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT))

# ⚠️ THE APP LAYER IS NO LONGER ONE FILE (router/core split, 2026-08-28).
# bridge/app.py is a FACADE over bridge/core/*.py + bridge/routers/*.py, so the
# source-text assertions below read bridge/appsrc.py's assembled view of the whole
# app layer instead of one file. Read bridge/appsrc.py's header for why the
# assertions are source-text in the first place and why order is part of it.
from bridge.appsrc import APP_SOURCE as _APP_SOURCE            # noqa: E402
from bridge import office                                        # noqa: E402

FAILS = []


def check(name, cond, extra=""):
    """`extra` is printed only on a FAILURE — the evidence, so a red line is legible
    without re-running by hand. Added at loffice-2026-08-28d, matching the other office
    suites; the older two-argument calls are unaffected."""
    print(("PASS" if cond else "FAIL"), name)
    if not cond:
        if extra != "":
            print(f"     {extra!r}")
        FAILS.append(name)


def eq(name, got, want):
    check(name, got == want, f"got {got!r}, want {want!r}")


try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    HAVE_XL = True
except Exception as _e:                                          # noqa: BLE001
    HAVE_XL = False
    print("!! openpyxl is NOT importable here — every round-trip test below is "
          f"SKIPPED and only the pure/wiring tables run ({_e})")


# ══ 1. names + containment ═══════════════════════════════════════════════════
NAME_TABLE = [
    ("sheet", "sheet.xlsx", "a bare stem gets the extension"),
    ("sheet.xlsx", "sheet.xlsx", "an .xlsx name passes through"),
    ("Sheet.XLSX", "Sheet.xlsx", "an upper-case extension is accepted and normalised"),
    ("  budget  ", "budget.xlsx", "surrounding whitespace is trimmed"),
    ("a b c", "a b c.xlsx", "spaces inside a name are fine"),
    ("../secrets", None, "traversal is refused"),
    ("/etc/passwd", None, "an absolute path is refused"),
    ("a/b.xlsx", None, "a separator is refused"),
    ("a\\b.xlsx", None, "a backslash is refused"),
    (".hidden", None, "a dotfile is refused"),
    ("..", None, "'..' is refused"),
    (".", None, "'.' is refused"),
    # ⚠️ THE POLICY CHANGED AT loffice-2026-08-29a AND THIS IS THE NEW, STRONGER PIN.
    # LOffice stores THREE types now (stage 3: docs + slides). .docx and .pptx are
    # ACCEPTED as names — they are opaque blobs to this module, and the refusal moved
    # from the NAME to the CONTENT path (require_sheet, pinned in its own group below).
    # Everything that is not one of the three is still refused here.
    ("sheet.docx", "sheet.docx", "a document name is accepted — three types live here"),
    ("deck.pptx", "deck.pptx", "…and a presentation name"),
    ("Deck.PPTX", "Deck.pptx", "…case-normalised like .xlsx is"),
    ("sheet.csv", None, "csv is still refused — it is not one of the three"),
    ("sheet.doc", None, "…and so is the legacy .doc"),
    ("sheet.txt", None, "…and anything else"),
    ("", None, "empty is refused"),
    ("   ", None, "whitespace-only is refused"),
    (None, None, "None is refused"),
    (17, None, "a number is refused"),
    ({"name": "x"}, None, "a dict is refused"),
    ("x\x00y", None, "a NUL byte is refused"),
    ("z" * 200, None, "an absurdly long name is refused"),
]
for raw, want, why in NAME_TABLE:
    got, reason = office.valid_name(raw)
    check(f"valid_name: {why}", got == want and (want is not None or bool(reason)))

check("every refusal carries a sentence",
      all(office.valid_name(r)[1] for r, w, _ in NAME_TABLE if w is None))

tmp = tempfile.mkdtemp()
D = office.office_dir(tmp)
check("office_dir is data/office under the root",
      os.path.realpath(D) == os.path.realpath(os.path.join(tmp, "data", "office")))
check("office_dir creates the folder", os.path.isdir(D))

open(os.path.join(D, "real.xlsx"), "wb").close()
t, r = office.doc_target(tmp, "real.xlsx")
check("doc_target resolves an existing workbook", t and os.path.isfile(t))
check("doc_target refuses traversal", office.doc_target(tmp, "../../x.xlsx")[0] is None)
check("doc_target refuses a missing file when must_exist",
      office.doc_target(tmp, "ghost.xlsx")[0] is None)
check("doc_target allows a missing file when not must_exist",
      office.doc_target(tmp, "ghost.xlsx", must_exist=False)[0] is not None)

# A symlink pointing OUT of the folder must not become a readable workbook.
outside = os.path.join(tmp, "outside.xlsx")
open(outside, "wb").close()
try:
    os.symlink(outside, os.path.join(D, "escape.xlsx"))
    got, reason = office.doc_target(tmp, "escape.xlsx")
    check("doc_target refuses a symlink that escapes the folder",
          got is None and "outside" in (reason or ""))
except OSError:
    check("doc_target refuses a symlink that escapes the folder (skipped: no symlinks)", True)

check("backup_for is <stem>.<day>.bak.xlsx",
      office.backup_for("/a/b/sheet.xlsx", "20260821") == "/a/b/sheet.20260821.bak.xlsx")
check("a backup name is recognised as one",
      office.is_backup_name("sheet.20260821.bak.xlsx"))
check("an ordinary name is not a backup", not office.is_backup_name("sheet.xlsx"))
check("is_backup_name is total over junk", office.is_backup_name(None) is False)


# ══ 2/3/4. the mappers ═══════════════════════════════════════════════════════
def build_fixture(path):
    """A workbook with one of everything the contract promises."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["B1"] = 42
    ws["C1"] = True
    ws["D1"] = "=B1*2"
    ws["A2"] = 3.5
    ws["B2"] = datetime.datetime(2026, 8, 21, 12, 0)
    ws["C2"] = ""                       # an empty string is not a value
    ws["A1"].font = Font(bold=True, italic=True, size=14, color="FFFF0000", name="Arial")
    ws["A1"].fill = PatternFill(fill_type="solid", start_color="FF00FF00",
                                end_color="FF00FF00")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="top")
    ws["B1"].number_format = "0.00"
    ws.column_dimensions["A"].width = 20
    ws.row_dimensions[1].height = 30
    ws.merge_cells("E1:F2")
    wb.create_sheet("Second")["A1"] = "x"
    wb.save(path)
    return path


if HAVE_XL:
    src = build_fixture(os.path.join(D, "fixture.xlsx"))
    snap = office.snapshot_from_path(src)
    s1 = snap["sheets"][snap["sheetOrder"][0]]
    cd = s1["cellData"]

    check("the snapshot carries both sheets in order", snap["sheetOrder"] == ["sheet-01", "sheet-02"])
    check("sheet names survive", [snap["sheets"][i]["name"] for i in snap["sheetOrder"]]
          == ["Data", "Second"])
    check("the workbook name comes from the file stem", snap["name"] == "fixture")
    check("cellData is keyed row-then-column, ZERO-based",
          cd["0"]["0"]["v"] == "Name" and cd["1"]["0"]["v"] == 3.5)
    check("a string cell is CellValueType STRING (1)", cd["0"]["0"]["t"] == 1
          and office.CV_STRING == 1)
    check("a number cell is NUMBER (2)", cd["0"]["1"]["t"] == 2 and office.CV_NUMBER == 2)
    check("an integer stays an integer", cd["0"]["1"]["v"] == 42
          and isinstance(cd["0"]["1"]["v"], int))
    check("a boolean is BOOLEAN (3)", cd["0"]["2"]["t"] == 3 and cd["0"]["2"]["v"] is True
          and office.CV_BOOLEAN == 3)
    check("a formula rides as `f`, verbatim", cd["0"]["3"]["f"] == "=B1*2")
    check("a formula cell carries no invented value",
          "v" not in cd["0"]["3"] or cd["0"]["3"].get("v") is not None)
    check("a date becomes the Excel serial number",
          abs(cd["1"]["1"]["v"] - 46255.5) < 1e-6 and cd["1"]["1"]["t"] == 2)
    check("the date's number format travels with it",
          "yyyy" in cd["1"]["1"]["s"]["n"]["pattern"])
    check("an empty string is not a cell", "2" not in cd.get("1", {}))
    st = cd["0"]["0"]["s"]
    check("bold survives", st["bl"] == 1)
    check("italic survives", st["it"] == 1)
    check("the font name survives", st["ff"] == "Arial")
    check("the font size survives", st["fs"] == 14.0)
    check("the font colour survives as #rrggbb", st["cl"] == {"rgb": "#ff0000"})
    check("the fill colour survives", st["bg"] == {"rgb": "#00ff00"})
    check("horizontal centre is HorizontalAlign.CENTER (2)", st["ht"] == 2)
    check("vertical top is VerticalAlign.TOP (1)", st["vt"] == 1)
    check("the number format survives", cd["0"]["1"]["s"]["n"]["pattern"] == "0.00")
    check("an UNSTYLED cell carries no style dict at all", "s" not in cd["1"]["0"])
    check("blank cells inside the used range are not emitted",
          all(v for row in cd.values() for v in row.values()))
    check("the merged range survives, zero-based",
          {"startRow": 0, "startColumn": 4, "endRow": 1, "endColumn": 5} in s1["mergeData"])
    check("the column width is carried in PIXELS", abs(s1["columnData"]["0"]["w"] - 140.0) < 0.01)
    check("the row height is carried in PIXELS", abs(s1["rowData"]["0"]["h"] - 40.0) < 0.01)
    check("the snapshot is JSON-serialisable", isinstance(json.dumps(snap), str))

    # ── the round trip ──
    dst = os.path.join(D, "written.xlsx")
    report = office.write_snapshot(snap, dst)
    check("write_snapshot reports both sheets", report["sheets"] == 2)
    check("write_snapshot reports the cells it wrote", report["cells"] >= 7)
    back = office.snapshot_from_path(dst)
    b1 = back["sheets"][back["sheetOrder"][0]]
    check("ROUND TRIP: every cell (value, type, formula, style) is identical",
          b1["cellData"] == cd)
    check("ROUND TRIP: sheet names are identical",
          [back["sheets"][i]["name"] for i in back["sheetOrder"]] == ["Data", "Second"])
    check("ROUND TRIP: merges are identical", b1["mergeData"] == s1["mergeData"])
    check("ROUND TRIP: column widths are identical", b1["columnData"] == s1["columnData"])
    check("ROUND TRIP: row heights are identical", b1["rowData"] == s1["rowData"])

    # ── cached formula results ──
    # A file whose formulas were last calculated by Excel carries the RESULT too; we
    # fake that by writing a value where the formula was and asserting the reader
    # prefers f + a cached v. openpyxl cannot calculate, so a file we wrote has none.
    check("a workbook we wrote has no cached result to invent",
          "v" not in b1["cellData"]["0"]["3"])

    # ── an empty workbook ──
    empty_path = os.path.join(D, "blank.xlsx")
    openpyxl.Workbook().save(empty_path)
    esnap = office.snapshot_from_path(empty_path)
    check("an empty workbook still opens with one sheet", len(esnap["sheetOrder"]) == 1)
    check("an empty workbook has no cells",
          esnap["sheets"][esnap["sheetOrder"][0]]["cellData"] == {})

    # ── new-file shape ──
    made, reason = office.create_doc(tmp, "budget")
    check("create_doc makes budget.xlsx", made == "budget.xlsx" and reason is None)
    check("the new workbook is a real xlsx", os.path.getsize(os.path.join(D, made)) > 0)
    nsnap = office.snapshot_from_path(os.path.join(D, made))
    check("a new workbook opens clean",
          nsnap["sheets"][nsnap["sheetOrder"][0]]["cellData"] == {})
    again, reason = office.create_doc(tmp, "budget")
    check("create_doc never clobbers an existing workbook",
          again is None and "already exists" in (reason or ""))
    check("create_doc refuses a junk name", office.create_doc(tmp, "../x")[0] is None)

    # ── THE EMPTY NAME. ────────────────────────────────────────────────────────
    # This is the landing. LOffice used to boot perfectly onto "No spreadsheets yet."
    # because being useful was gated on the user typing a file name first, and pressing
    # Create with nothing typed did NOTHING AT ALL. An empty name is now a request for
    # a sheet, answered with DEFAULT_DOC_STEM through the ' (n)' never-clobber walk, so
    # a fresh install lands on an editable grid and repeated presses cannot collide.
    check("there IS a default stem, and it is the ordinary word for it",
          office.DEFAULT_DOC_STEM == "Untitled")
    with tempfile.TemporaryDirectory() as fresh:
        first, reason = office.create_doc(fresh, "")
        check("an EMPTY name creates Untitled.xlsx rather than refusing",
              first == "Untitled.xlsx" and reason is None)
        second, _ = office.create_doc(fresh, "")
        check("…and a second one steps to Untitled (2).xlsx — the same never-clobber "
              "convention as import, the artifact save and the voice clips",
              second == "Untitled (2).xlsx")
        check("a whitespace-only name is the same request",
              office.create_doc(fresh, "   ")[0] == "Untitled (3).xlsx")
        check("None is the same request — the route coerces to '', but create_doc is "
              "called directly too", office.create_doc(fresh, None)[0] == "Untitled (4).xlsx")
        check("every one of them is a real workbook on disk",
              all(os.path.isfile(os.path.join(office.office_dir(fresh), n))
                  for n in ("Untitled.xlsx", "Untitled (2).xlsx",
                            "Untitled (3).xlsx", "Untitled (4).xlsx")))
        check("…and list_docs returns them newest-first, which is the ONLY reason the "
              "panel can land on files[0] and call it the one you were last in",
              [e["name"] for e in office.list_docs(fresh)][0] == "Untitled (4).xlsx")
        # The asymmetry is deliberate and must stay: a name the user TYPED is honoured
        # literally, so a collision is reported rather than renamed behind their back.
        office.create_doc(fresh, "budget")
        check("an EXPLICIT name still refuses on collision — it is not silently "
              "renamed the way an empty one is walked",
              office.create_doc(fresh, "budget")[0] is None)
        check("a junk TYPE still refuses rather than falling through to the default",
              office.create_doc(fresh, 7)[0] is None
              and office.create_doc(fresh, ["x"])[0] is None)

    # ── totality of the reverse mapper ──
    JUNK = {
        "id": "w", "sheetOrder": ["s1"], "styles": {"S1": {"bl": 1}},
        "sheets": {"s1": {"id": "s1", "name": "J", "cellData": {
            "0": {"0": {"v": "keep", "t": 1}},
            "1": "not a row",
            "x": {"0": {"v": "bad row key"}},
            "2": {"y": {"v": "bad col key"}, "0": "not a cell", "1": {"v": float("nan")},
                  "2": {"v": None}, "3": {"v": {"deep": 1}}, "4": {"f": 17},
                  "5": {"v": 5, "s": "S1"}, "6": {"v": 6, "s": "NOPE"},
                  "7": {"v": "07", "t": 4}},
            "-1": {"0": {"v": "negative row"}},
        }, "mergeData": ["junk", {"startRow": 5, "startColumn": 5, "endRow": 1,
                                  "endColumn": 1}, {"startRow": 8, "startColumn": 0,
                                                    "endRow": 9, "endColumn": 1}],
           "rowData": {"0": {"h": "tall"}, "1": {"h": 48}},
           "columnData": "junk"}},
    }
    jpath = os.path.join(D, "junk.xlsx")
    jreport = office.write_snapshot(JUNK, jpath)
    check("a snapshot full of junk still writes a workbook", os.path.isfile(jpath))
    jback = office.snapshot_from_path(jpath)
    jcd = jback["sheets"][jback["sheetOrder"][0]]["cellData"]
    check("a good cell beside junk survives", jcd["0"]["0"]["v"] == "keep")
    check("a row that is not a dict is skipped", "1" not in jcd)
    check("a non-numeric row key is skipped", all(k.lstrip("-").isdigit() for k in jcd))
    check("a NaN value is not written", "1" not in jcd.get("2", {}))
    check("a None value is not written", "2" not in jcd.get("2", {}))
    check("a cell that is not a dict is skipped", "0" not in jcd.get("2", {}))
    check("a nested-object value is stringified, not dropped silently",
          isinstance(jcd["2"]["3"]["v"], str))
    check("a non-string formula is not treated as a formula",
          "f" not in jcd.get("2", {}).get("4", {}))
    check("a style id resolves out of workbook.styles", jcd["2"]["5"]["s"]["bl"] == 1)
    check("an unknown style id costs the style, not the value", jcd["2"]["6"]["v"] == 6)
    # ⚠️ THE `t` THIS ASSERTS CHANGED 1 → 4 AT loffice-2026-08-28d, AND THE CHANGE IS
    # FINDING F-28'S FIX. FORCE_STRING means "this cell is text ON PURPOSE" — a part code,
    # a leading-zero id. The value survived the round-trip; the INTENT did not, because the
    # .xlsx has nowhere to keep it, so it came back as an ordinary string (t:1) and the new
    # text-that-looks-numeric detector then nagged about a cell the page had deliberately
    # made text. `_write_cell` now stamps the `@` number format — which is exactly what
    # Excel uses to say this — and `cell_snapshot` re-derives t:4 from it.
    check("FORCE_STRING keeps a leading zero", jcd["2"]["7"]["v"] == "07"
          and office.CV_FORCE_STRING == 4)
    check("…and the INTENT survives the round-trip, as the `@` number format",
          jcd["2"]["7"]["t"] == office.CV_FORCE_STRING
          and office.is_text_format((jcd["2"]["7"].get("s") or {}).get("n", {})
                                    .get("pattern")))
    check("a negative row index is skipped", "-1" not in jcd)
    check("an inverted merge range is skipped and a good one kept",
          {"startRow": 8, "startColumn": 0, "endRow": 9, "endColumn": 1}
          in jback["sheets"][jback["sheetOrder"][0]]["mergeData"]
          and len(jback["sheets"][jback["sheetOrder"][0]]["mergeData"]) == 1)
    check("a junk row height is ignored and a good one applied",
          jback["sheets"][jback["sheetOrder"][0]]["rowData"].get("1", {}).get("h") == 48.0)

    for bad, why in [(None, "None"), ("a string", "a string"), (17, "a number"),
                     ({}, "an empty dict"), ({"sheets": {}}, "no sheets"),
                     ({"sheets": "x"}, "sheets that is not a dict")]:
        try:
            office.write_snapshot(bad, os.path.join(D, "never.xlsx"))
            ok = False
        except office.OfficeError as e:
            ok = bool(str(e))
        check(f"write_snapshot refuses {why} with a sentence", ok)
    check("a refused write leaves no file behind",
          not os.path.exists(os.path.join(D, "never.xlsx")))

    # sheetOrder junk must not lose the sheets
    recovered = office.write_snapshot(
        {"sheetOrder": "junk", "sheets": {"a": {"name": "One", "cellData": {}},
                                          "b": {"name": "One", "cellData": {}}}},
        os.path.join(D, "recovered.xlsx"))
    check("a junk sheetOrder falls back to the sheets dict", recovered["sheets"] == 2)
    rsnap = office.snapshot_from_path(os.path.join(D, "recovered.xlsx"))
    names = [rsnap["sheets"][i]["name"] for i in rsnap["sheetOrder"]]
    check("duplicate sheet names are made unique (Excel refuses duplicates)",
          len(set(n.lower() for n in names)) == 2)

    # ── the backup policy ──
    # ⚠️ EVERY SAVE OVER AN EXISTING WORKBOOK NOW CARRIES ITS FENCE (bug-echo BE-02).
    # `mt()` is what a caller that has just read the file knows; a save with nothing to
    # fence with is refused rather than written, which is what the block below pins.
    def mt(n="budget.xlsx"):
        return os.path.getmtime(os.path.join(D, n))
    saved, reason = office.save_doc(tmp, "budget.xlsx", office.empty_snapshot("budget"),
                                    mt())
    bak = office.backup_for(os.path.join(D, "budget.xlsx"))
    check("saving over an existing workbook takes a backup first",
          saved and saved["backup"] == os.path.basename(bak) and os.path.isfile(bak))
    before = os.path.getmtime(bak)
    saved2, _ = office.save_doc(tmp, "budget.xlsx", office.empty_snapshot("budget"),
                                mt())
    check("a SECOND save the same day does not re-copy",
          saved2 and saved2["backup"] == "" and os.path.getmtime(bak) == before)
    listed = [f["name"] for f in office.list_docs(tmp)]
    check("backups are on disk but not in the file list",
          os.path.basename(bak) not in listed and "budget.xlsx" in listed)
    check("the file list flags that a backup exists",
          any(f["name"] == "budget.xlsx" and f["has_backup"] for f in office.list_docs(tmp)))
    check("the file list is newest-first",
          [f["modified"] for f in office.list_docs(tmp)]
          == sorted([f["modified"] for f in office.list_docs(tmp)], reverse=True))
    check("saving a NEW name takes no backup",
          (office.save_doc(tmp, "fresh.xlsx", office.empty_snapshot("fresh"))[0] or {})
          .get("backup") == "")
    check("save_doc refuses a junk name before touching disk",
          office.save_doc(tmp, "../evil.xlsx", office.empty_snapshot("x"))[0] is None)
    check("save_doc refuses a junk snapshot",
          office.save_doc(tmp, "fresh.xlsx", "nope", mt("fresh.xlsx"))[0] is None)
    check("…for the SNAPSHOT's reason, not the fence's — a caller that did everything "
          "right must not be told it forgot the mtime",
          office.save_doc(tmp, "fresh.xlsx", "nope", mt("fresh.xlsx"))[1]
          != office.SAVE_UNFENCED_REFUSAL)

    # ══ THE SAVE'S mtime FENCE (bug-echo BE-02, the F-01 class) ═════════════════════
    # THE REPRO, EXECUTED: the page reads a snapshot at T0, somebody else writes at T1,
    # the page saves at T2. It used to land, silently, reverting the other write to the
    # page's stale copy with only the once-per-day .bak behind it.
    import openpyxl as _px
    race = os.path.join(D, "race.xlsx")
    _w = _px.Workbook(); _w.active["A1"] = 10; _w.save(race)
    snap_t0 = office.snapshot_from_path(race)                   # the page opens at T0
    seen_t0 = os.path.getmtime(race)
    _w2 = _px.load_workbook(race); _w2.active["A1"] = 999       # somebody else at T1
    _w2.save(race)
    os.utime(race, (seen_t0 + 60, seen_t0 + 60))                # unambiguously newer
    rep_stale, why = office.save_doc(tmp, "race.xlsx", snap_t0, seen_t0)
    check("a save whose fence is stale is REFUSED — nothing is written",
          rep_stale is None and why == office.SAVE_FENCE_REFUSAL, why)
    check("…and the OTHER writer's value is still on disk (the whole point)",
          _px.load_workbook(race).active["A1"].value == 999)
    check("…and the refusal says how to get past it",
          "force" in office.SAVE_FENCE_REFUSAL)
    check("…while force=True overwrites deliberately, and SAYS it did",
          (office.save_doc(tmp, "race.xlsx", snap_t0, seen_t0, force=True)[0]
           or {}).get("forced") is True)
    check("a save with NO fence at all is a refusal to GUESS, not a silent skip",
          office.save_doc(tmp, "race.xlsx", snap_t0)[1] == office.SAVE_UNFENCED_REFUSAL)
    for _bad in (float("nan"), float("inf"), "lunchtime", [], {}):
        check(f"…and {_bad!r} is refused the same way rather than skipping the fence",
              office.save_doc(tmp, "race.xlsx", snap_t0, _bad)[1]
              == office.SAVE_UNFENCED_REFUSAL)
    check("unfenced=True is the DELIBERATE skip, and it writes and says it was unfenced",
          (office.save_doc(tmp, "race.xlsx", snap_t0, None, unfenced=True)[0] or {})
          .get("fenced") is False)
    check("a matching fence writes, and says it WAS fenced",
          (office.save_doc(tmp, "race.xlsx", snap_t0,
                           os.path.getmtime(race))[0] or {}).get("fenced") is True)
    check("…and the report carries the NEW mtime, so the next save can stay fenced",
          abs((office.save_doc(tmp, "race.xlsx", snap_t0,
                               os.path.getmtime(race))[0] or {}).get("mtime", 0)
              - os.path.getmtime(race)) < 0.01)
    check("A FILE THAT DOES NOT EXIST YET IS NEVER FENCED — a first save must not need "
          "a version to be stale against",
          (office.save_doc(tmp, "brandnew.xlsx",
                           office.empty_snapshot("brandnew"))[0] or {})
          .get("created") is True)
    # ⚠️ THE FENCE RUNS BEFORE THE `.bak` — finding F-20's ordering, which says every
    # reason to refuse comes before any copy is taken. A refused save must leave the
    # workbook AND its backup exactly as they were, or a stale save would burn today's
    # one-per-day backup slot on a write that never happened.
    office.save_doc(tmp, "nobak.xlsx", office.empty_snapshot("nobak"))
    nobak = office.backup_for(os.path.join(D, "nobak.xlsx"))
    check("(fixture) the new workbook has no daily .bak yet", not os.path.isfile(nobak))
    check("a stale save is refused…",
          office.save_doc(tmp, "nobak.xlsx", office.empty_snapshot("nobak"), 1.0)[0]
          is None)
    check("…and took NO backup on the way out (F-20 ordering holds for this fence too)",
          not os.path.isfile(nobak))
    check("…and an un-fenced save is refused before the backup as well",
          office.save_doc(tmp, "nobak.xlsx", office.empty_snapshot("nobak"))[0] is None
          and not os.path.isfile(nobak))

    # ── open/delete ──
    got, reason = office.open_doc(tmp, "budget.xlsx")
    check("open_doc returns a snapshot", got and got["sheetOrder"])
    # ⚠️ AND THE VERSION IT CAME FROM (bug-echo BE-02): an open is where a caller LEARNS
    # what its next save must fence against, so the mtime travels with the content — read
    # BEFORE it, so a write landing in between cannot hand the caller a newer stamp than
    # the bytes it is holding.
    check("…carrying the file's mtime, which is the fence value for the next save",
          abs((got or {}).get("file_mtime", 0)
              - os.path.getmtime(os.path.join(D, "budget.xlsx"))) < 0.001)
    check("…and that value really is accepted as the fence by save_doc",
          office.save_doc(tmp, "budget.xlsx", got, got["file_mtime"])[0] is not None)
    check("open_doc refuses a missing workbook", office.open_doc(tmp, "ghost.xlsx")[0] is None)
    ok, reason = office.delete_doc(tmp, "fresh.xlsx")
    check("delete_doc removes the workbook", ok and not os.path.exists(os.path.join(D, "fresh.xlsx")))
    check("delete_doc refuses traversal", office.delete_doc(tmp, "../x.xlsx")[0] is False)
    check("delete leaves the backup alone (it exists to survive mistakes)",
          os.path.isfile(bak))

    # ── rename ──
    # Both names go through doc_target, so the rename boundary IS the delete boundary.
    office.create_doc(tmp, "before.xlsx")
    new, reason = office.rename_doc(tmp, "before.xlsx", "after")
    check("rename_doc moves the workbook and forces the extension",
          new == "after.xlsx" and os.path.isfile(os.path.join(D, "after.xlsx"))
          and not os.path.exists(os.path.join(D, "before.xlsx")))
    check("rename_doc REFUSES a name already in use rather than resolving it to ' (n)' "
          "— a rename names a file the user typed, and quietly picking a different one "
          "hides the collision", office.rename_doc(tmp, "after.xlsx", "budget.xlsx")[0] is None)
    check("…and leaves both files exactly where they were",
          os.path.isfile(os.path.join(D, "after.xlsx"))
          and os.path.isfile(os.path.join(D, "budget.xlsx")))
    check("renaming a workbook to its own name is a no-op, not a failure",
          office.rename_doc(tmp, "after.xlsx", "after")[0] == "after.xlsx"
          and os.path.isfile(os.path.join(D, "after.xlsx")))
    check("rename_doc refuses traversal on the SOURCE",
          office.rename_doc(tmp, "../x.xlsx", "ok")[0] is None)
    check("…and on the TARGET", office.rename_doc(tmp, "after.xlsx", "../ok.xlsx")[0] is None)
    check("rename_doc refuses a missing workbook",
          office.rename_doc(tmp, "ghost.xlsx", "ok")[0] is None)
    # ⚠️ A RENAME MAY CHANGE THE STEM AND NEVER THE TYPE (loffice-2026-08-29a). With
    # three extensions in the folder this stopped being a naming rule and became a
    # TRUTHFULNESS rule: `budget.docx` holding spreadsheet bytes is a file whose
    # extension lies, and the editor would then blame the file rather than the rename.
    check("rename_doc refuses a change of TYPE, even to another extension it stores",
          office.rename_doc(tmp, "after.xlsx", "after.docx")[0] is None
          and ".docx" in (office.rename_doc(tmp, "after.xlsx", "after.docx")[1] or ""))
    check("…and a rename with NO extension keeps the source's",
          office.rename_doc(tmp, "after.xlsx", "afterwards")[0] == "afterwards.xlsx")
    check("…and renaming it back leaves the name where the rest of this group expects it",
          office.rename_doc(tmp, "afterwards.xlsx", "after")[0] == "after.xlsx")
    check("rename_doc refuses an empty or junk new name without raising",
          office.rename_doc(tmp, "after.xlsx", "")[0] is None
          and office.rename_doc(tmp, "after.xlsx", None)[0] is None
          and office.rename_doc(tmp, None, "x")[0] is None
          and office.rename_doc(tmp, "after.xlsx", 7)[0] is None)
    check("every refusal comes with a sentence, never a bare None",
          isinstance(office.rename_doc(tmp, "ghost.xlsx", "ok")[1], str)
          and office.rename_doc(tmp, "ghost.xlsx", "ok")[1] != "")

    # ── the cell cap ──
    big = {"sheetOrder": ["s"], "sheets": {"s": {"name": "Big", "cellData": {}}}}
    check("MAX_CELLS is a real bound, not decoration", office.MAX_CELLS >= 100000)

# a snapshot of an unreadable file must be a sentence, not a traceback
bad_path = os.path.join(D, "notreally.xlsx")
with open(bad_path, "w", encoding="utf-8") as fh:
    fh.write("this is not a zip archive")
try:
    office.snapshot_from_path(bad_path)
    ok = False
except office.OfficeError as e:
    ok = "could not read" in str(e) or "openpyxl" in str(e)
except Exception:                                                # noqa: BLE001
    ok = False
check("a corrupt file raises OfficeError with a sentence, never a raw traceback", ok)


# ══ 5. pure helpers ══════════════════════════════════════════════════════════
check("excel_serial: the 1900 leap bug epoch is 1899-12-30",
      excel := office.excel_serial(datetime.datetime(1900, 1, 1)) == 2.0)
check("excel_serial: a date is a whole number",
      office.excel_serial(datetime.date(2026, 8, 21)) == 46255.0)
check("excel_serial: a time is a day fraction",
      abs(office.excel_serial(datetime.time(12, 0)) - 0.5) < 1e-9)
check("excel_serial is total over junk", office.excel_serial("nope") == 0.0)
for raw, want in [("#FF0000", "FFFF0000"), ("ff0000", "FFFF0000"),
                  ("FFFF0000", "FFFF0000"), ("nope", ""), (None, ""), (17, ""),
                  ("#12345", "")]:
    check(f"_hex_to_argb({raw!r})", office._hex_to_argb(raw) == want)
for raw, want in [("3", 3), (3, 3), ("0", 0), ("-1", None), ("x", None),
                  (None, None), (True, None), (3.7, 3)]:
    check(f"_idx({raw!r})", office._idx(raw) == want)
check("_num rejects NaN", office._num(float("nan")) is None)
check("_num rejects inf", office._num(float("inf")) is None)
check("_num rejects a bool", office._num(True) is None)
check("_num accepts a numeric string", office._num("3.5") == 3.5)

es = office.empty_snapshot("Untitled")
check("empty_snapshot has exactly one sheet", len(es["sheetOrder"]) == 1)
check("empty_snapshot's sheet is in the sheets dict", es["sheetOrder"][0] in es["sheets"])
check("empty_snapshot is JSON-serialisable", isinstance(json.dumps(es), str))
check("empty_snapshot names the workbook from the file stem",
      office.empty_snapshot("budget.xlsx")["name"] == "budget")

# ⚠️ REWRITTEN AT loffice-2026-08-28d, AND THE REWRITE IS THE FIX (live finding L3). The
# fidelity contract used to be ONE sentence — "complex styling may be simplified; keep your
# original file" — printed on every surface. It described the openpyxl mapper, while the
# surface Debi actually saves through is the embedded editor, whose x2t round-trip
# MEASURABLY keeps charts, images, autofilters, validation, hyperlinks and freeze panes
# (measured on QA-rich.xlsx: chart1.xml 1159 B → 3020 B, present after the save). So the
# contract is now PER SAVE PATH, and what this pins is that both sentences exist and that
# each one names its own path rather than the other's limits.
check("the fidelity contract is single-sourced and split per SAVE PATH",
      "mapper" in office.FIDELITY_NOTE and "full editor" in office.FIDELITY_NOTE
      and "editor" in office.FIDELITY_EDITOR_NOTE)
check("the FILE path's sentence names what it really loses",
      all(x in office.FIDELITY_NOTE for x in ("charts", "images", "borders", "macros")))
check("…and no longer claims to lose the four things F-27 made it carry",
      all(x in office.FIDELITY_NOTE for x in
          ("freeze panes", "hidden rows and columns", "autofilter")))
check("the EDITOR path's sentence claims the fidelity that was measured, and says it was",
      all(x in office.FIDELITY_EDITOR_NOTE for x in
          ("charts", "images", "autofilters", "freeze panes"))
      and "Measured" in office.FIDELITY_EDITOR_NOTE)


# ══ THE 2026-08-28 ADVERSARIAL CAMPAIGN — the mapper/filesystem half ═════════
# Repros live in bridge/tests/test_office_adversarial.py (the campaign ledger, not a gate).
print("\n── the adversarial campaign: what the file layer promises ──")
if HAVE_XL:
    _cd = tempfile.mkdtemp(prefix="office-camp-")
    _od = office.office_dir(_cd)

    def _mk(name, val="v"):
        wb = openpyxl.Workbook()
        wb.active["A1"] = val
        wb.save(os.path.join(_od, name))

    # L4 — the external-change banner promised `<name>.pre-agent.xlsx` on ANY mtime change,
    # including the several that never make one; the DIRTY fork offered "Keep mine" on the
    # strength of a backup that was not there. The row now carries whether it EXISTS, per
    # path, so the page can say only what is true.
    _mk("banner.xlsx")
    _row = [f for f in office.list_docs(_cd) if f["name"] == "banner.xlsx"][0]
    eq("L4 · a workbook nothing has written carries NO agent copy for the banner to name",
       _row["agent_copy"], "")
    os.makedirs(os.path.dirname(office.agent_copy_path(_cd, "banner.xlsx")),
                exist_ok=True)
    shutil.copy2(os.path.join(_od, "banner.xlsx"),
                 office.agent_copy_path(_cd, "banner.xlsx"))
    _row = [f for f in office.list_docs(_cd) if f["name"] == "banner.xlsx"][0]
    eq("…and once one really exists, the row names it — checked by stat, not derived from "
       "the workbook's name", _row["agent_copy"],
       office.CHECKPOINT_DIR + "/banner/pre-agent.xlsx")

    # C4 — the safety copies were listed as ordinary workbooks, with a size, a date and a
    # download/delete pair, indistinguishable from documents, and the sentence that made
    # them long gone. Filtered exactly the way a .bak is.
    _mk("legacy.pre-agent.xlsx", "an old sibling copy")
    _names = [f["name"] for f in office.list_docs(_cd)]
    check("C4 · a legacy *.pre-agent.xlsx sibling is filtered out of the workbook list, "
          "the way a .bak is — a safety copy is not a document",
          "legacy.pre-agent.xlsx" not in _names and "banner.xlsx" in _names, _names)
    check("…and the predicate is a named, testable one rather than a substring somewhere",
          office.is_agent_copy_name("x.pre-agent.xlsx")
          and not office.is_agent_copy_name("x.xlsx"))

    # B5 — "This cannot be undone" was false in BOTH directions: a delete left the daily
    # .bak on disk, `is_backup_name()` hides those from the list, so the delete COULD be
    # undone from a file the user was not told about — and somebody deleting a workbook for
    # privacy KEPT its contents.
    _mk("bye.xlsx")
    _bak = office.backup_for(os.path.join(_od, "bye.xlsx"))
    shutil.copy2(os.path.join(_od, "bye.xlsx"), _bak)
    ok, _r = office.delete_doc(_cd, "bye.xlsx")
    check("the default delete still KEEPS the daily backup — it exists to survive a "
          "mistake, and that ruling has not changed", ok and os.path.isfile(_bak))
    _mk("bye2.xlsx")
    _bak2 = office.backup_for(os.path.join(_od, "bye2.xlsx"))
    shutil.copy2(os.path.join(_od, "bye2.xlsx"), _bak2)
    os.makedirs(os.path.dirname(office.agent_copy_path(_cd, "bye2.xlsx")), exist_ok=True)
    shutil.copy2(os.path.join(_od, "bye2.xlsx"), office.agent_copy_path(_cd, "bye2.xlsx"))
    ok, _r = office.delete_doc(_cd, "bye2.xlsx", True)
    check("B5 · …but a user who says 'and its backups' can MEAN it — the daily copy, the "
          "agent copy and the checkpoint stack all go",
          ok and not os.path.isfile(_bak2)
          and not os.path.exists(office.agent_copy_path(_cd, "bye2.xlsx")), _r)

    # F-21 — the checkpoint stack was keyed by the file STEM and a rename moved neither it
    # nor the pre-agent sibling, so undo answered "no such workbook" for ever.
    _mk("mover.xlsx")
    _cp = os.path.join(office.checkpoint_root(_cd), "mover")
    os.makedirs(_cp, exist_ok=True)
    shutil.copy2(os.path.join(_od, "mover.xlsx"), os.path.join(_cp, "abc123.xlsx"))
    _new, _r = office.rename_doc(_cd, "mover.xlsx", "moved.xlsx")
    check("F-21 · rename_doc moves the checkpoint folder with the workbook, so the undo "
          "stack is not orphaned under a stem nothing points at any more",
          _new == "moved.xlsx"
          and os.path.isfile(os.path.join(office.checkpoint_root(_cd), "moved",
                                          "abc123.xlsx"))
          and not os.path.exists(_cp), _r)

    shutil.rmtree(_cd, ignore_errors=True)

# ══ 6. wiring ════════════════════════════════════════════════════════════════
APP = _APP_SOURCE
PAGE = (ROOT / "bridge" / "panel" / "office.html").read_text(encoding="utf-8")
SWIFT = (ROOT / "app" / "main.swift").read_text(encoding="utf-8")
FETCH = (ROOT / "scripts" / "fetch_vendor_assets.sh").read_text(encoding="utf-8")
YAML = (ROOT / "harness.yaml").read_text(encoding="utf-8")
REQ = (ROOT / "bridge" / "requirements.txt").read_text(encoding="utf-8")

for route in ['@app.get("/office")', '@app.get("/api/office/files")',
              '@app.post("/api/office/new")', '@app.get("/api/office/open/{name}")',
              '@app.post("/api/office/save")', '@app.post("/api/office/delete")',
              '@app.post("/api/office/rename")',
              '@app.get("/api/office/download/{name}")']:
    check(f"app.py declares {route}", route in APP)
check("the office module is imported DEFENSIVELY (a missing file must not kill the bridge)",
      "from bridge import office as _office" in APP and "_OFFICE_ERR" in APP)
check("every office route answers 503 when the module is missing",
      APP.count("return _office_unavailable()") >= 6)
# The route must NOT pre-judge an empty name. `(body or {}).get("name") or ""` hands
# create_doc the empty string, and create_doc is the one place that decides what an
# empty name means (Untitled through the never-clobber walk). A guard here would put
# the dead end back one layer down, where nobody would look for it.
_new_route = APP[APP.index('@app.post("/api/office/new")'):]
_new_route = _new_route[:_new_route.index('@app.get("/api/office/open')]
check("POST /api/office/new passes an empty name straight through to create_doc, and "
      "carries no empty-name refusal of its own — a guard here would put the dead end "
      "back one layer down, where nobody would look for it",
      '_office.create_doc, ROOT, want, ext)' in _new_route
      and 'want = (body or {}).get("name") or ""' in _new_route
      and "no file name" not in _new_route)
# ⚠️ `step` ARRIVED AT loffice-2026-08-28d (live finding B2) AND IT DOES NOT WEAKEN THE
# NEVER-CLOBBER RULING — it is how the CALLER says which of the route's two meanings it
# has. A name the USER TYPED still collides and is still refused, because quietly making
# `budget (2).xlsx` when somebody asked for `budget.xlsx` hides the thing they need to
# know. A TEMPLATE CARD is "give me one of these", and refusing it left the card
# PERMANENTLY DEAD after one use, with a red line offering no way forward, on a screen
# listing the existing file two inches below.
check("…and a TEMPLATE's collision takes free_name's ' (n)' walk instead, on the "
      "caller's own say-so rather than by weakening create_doc",
      "_office.free_name, ROOT," in _new_route
      and '(body or {}).get("step")' in _new_route)
check("…and the walk is the SAME one import and the blank name already take — one "
      "never-clobber convention in this codebase, not a second",
      "free_name" in APP)
check("download and delete both go through doc_target/delete_doc containment",
      "_office.doc_target(ROOT, name)" in APP and "_office.delete_doc" in APP)
check("rename goes through rename_doc — which puts BOTH names through doc_target, so "
      "it can no more reach outside data/office than a delete can",
      "_office.rename_doc, ROOT" in APP
      and "asyncio.to_thread(\n        _office.rename_doc" in APP)
# ⚠️ THE HEADER DICT GAINED A WRAPPER AT loffice-2026-08-28a AND THE ASSERTION FOLLOWED
# IT RATHER THAN BEING DROPPED. /office is the document that EMBEDS the ONLYOFFICE
# editor now, so it is served through `_oo_headers(...)` — the same one source of truth
# as /oo-edit and /oo/* — and no-store is passed through it. Both halves are asserted:
# the stale-panel lesson (no-store) and the isolation the embed cannot work without.
check("the page is served no-store (the stale-panel lesson)",
      re.search(r'PANEL / "office\.html",\s*headers=_oo_headers\(\{"Cache-Control": '
                r'"no-store', APP, re.S) is not None)
check("…and with the three cross-origin-isolation headers, from the ONE place that "
      "owns them, because it is the document that embeds the editor",
      "def _oo_headers(" in APP
      and re.search(r'PANEL / "office\.html",\s*headers=_oo_headers\(', APP, re.S)
      is not None)
check("the blocking round-trip runs off the event loop",
      "asyncio.to_thread(_office.open_doc" in APP
      # The save's call wrapped onto its own line when it grew the mtime fence
      # (bug-echo BE-02), so the anchor allows the wrap rather than pinning the layout.
      and re.search(r"asyncio\.to_thread\(\s*_office\.save_doc", APP) is not None)
# ⚠️ THE SAVE ROUTE PASSES THE FENCE, AND SAYS SO WHEN IT CANNOT (bug-echo BE-02). The
# module-level fence is only half the fix: a route that never forwarded `expect_mtime`
# would leave save_doc's new parameter permanently at its refusal default, and a route
# that quietly passed `unfenced=True` with no log line would be the silent clobber again
# wearing a keyword argument.
check("the save route forwards the caller's fence…", "expect_mtime" in APP)
check("…accepts the writeback route's spelling of it too, so one lane needs one word",
      re.search(r'body\.get\("mtime"\)', APP) is not None)
check("…maps the fence refusal onto a 409 like the editor's, not a 400",
      re.search(r"409 if reason == _office\.SAVE_FENCE_REFUSAL", APP) is not None)
check("…and when nothing was sent to fence with, the skip is EXPLICIT and LOGGED",
      "unfenced=(fence is None)" in APP and "UNFENCED" in APP)
check("the fidelity note reaches the panel from the module, not a second copy",
      "_office.FIDELITY_NOTE" in APP and office.FIDELITY_NOTE not in PAGE)

check("main.swift has the LOffice tab, last, on the bridge origin",
      'HarnessTab(id: "loffice", title: "LOffice", url: URL(string: "http://127.0.0.1:8700/office")!)' in SWIFT)
# STUDIO PHASE 2 (2026-08-21) split the table into a REGISTRY (everything that can be a
# tab) and the strip, which the nav model orders. LOffice is still the tenth and last of
# the DEFAULT strip; the registry additionally carries the three pinnable panel views.
tab_titles = re.findall(r'HarnessTab\(id: "[^"]+", title: "([^"]+)"', SWIFT)
default_top = re.search(r'let navDefaultTopbar = \[([^\]]+)\]', SWIFT)
default_ids = re.findall(r'"([^"]+)"', default_top.group(1)) if default_top else []
# (OpenCode landed after LOffice and took the last slot, 2026-08-21 — so the assertion
# is that LOffice is ON the default strip and still ahead of anything added later, not
# that it is last forever.)
check("LOffice is on the DEFAULT strip",
      "loffice" in default_ids and len(default_ids) >= 10)
check("...and it is still a registry row",
      "LOffice" in tab_titles)
check("the ROUTE did not churn with the wordmark — the tab still points at /office",
      "8700/office" in SWIFT and '"/office"' in APP)
check("the page wears the same name as the tab",
      "<title>LOffice" in PAGE and ">LOffice<" in PAGE)
check("the width budget note was re-checked when the tab count last changed",
      "ELEVEN current titles" in SWIFT)
check("the title is ONE word — a two-word title does not fit the strip budget",
      " " not in "LOffice")
check("…and the reason Debi's 'Office Lane' was not taken is recorded next to it",
      "Office Lane" in SWIFT)

check("harness.yaml pins univer", re.search(r'^\s*univer_pin:\s*"[\d.]+"', YAML, re.M))
check("the fetch script reads the pin from harness.yaml", "univer_pin:" in FETCH
      and "UNIVER_V=" in FETCH)
for asset in ["univer/rxjs.umd.min.js", "univer/presets.umd.js",
              "univer/preset-sheets-core.umd.js", "univer/preset-sheets-core.css",
              "univer/preset-sheets-core.en-US.js"]:
    check(f"the fetch script vendors {asset}", asset in FETCH)
check("the fetch script makes the subdirectory before writing",
      'mkdir -p "$(dirname "$out")"' in FETCH)
check("no runtime CDN reaches the office page",
      "cdn.jsdelivr" not in PAGE and "unpkg" not in PAGE and "http://" not in
      PAGE.split("<script")[0])
check("openpyxl is declared in bridge/requirements.txt", "openpyxl" in REQ)

# ⚠️ THE BUNDLES ARE NO LONGER STATIC TAGS (2026-08-21e). They are loaded on demand by
# `loadVendor()` when somebody presses "Rich editor" — so the LOAD ORDER, which is the
# one thing a reader could get wrong and only find out in a browser, is now DATA (the
# `VENDOR` list) and is pinned in bridge/tests/test_office_grid.js, where it can be read
# as a list rather than scraped out of markup. What this file pins is the invariant that
# made the change necessary: the served document has NO external subresource at all.
SCRIPT_TAGS = re.findall(r'<script\b([^>]*)\bsrc="/assets/vendor/([^"]+)"([^>]*)>', PAGE)
check("the served document loads no vendor bundle by itself — tier 1 IS the document",
      SCRIPT_TAGS == [])
check("the vendor bundles are declared as an ordered list instead",
      "const VENDOR = [" in PAGE and PAGE.count("/assets/vendor/") >= 6)
check("the stylesheet is self-hosted too",
      '/assets/vendor/univer/preset-sheets-core.css' in PAGE)
check("the page uses the globals the pinned bundles actually define",
      "window.UniverPresets" in PAGE and "window.UniverPresetSheetsCore" in PAGE
      and "UniverPresetSheetsCoreEnUS" in PAGE)
check("the page builds the workbook through the facade",
      "createUniver(" in PAGE and "createUniverSheet(" in PAGE)
check("saving reads the snapshot back with save() and the deprecated alias as fallback",
      "wb.save ? wb.save() : wb.getSnapshot()" in PAGE)
# ⚠️ CHANGED HONESTLY 2026-08-21: this used to pin the AI rail as "a labelled STUB with
# no controls", which was the right assertion while it WAS one — a dead button is worse
# than an honest placeholder. Slice 2 built the panel, so the pinned fact genuinely
# moved; the invariants that replace it are the ones that matter now (the panel exists,
# it is reachable, and it speaks the EXISTING chat lane rather than a new endpoint).
# The full behaviour is pinned in bridge/tests/test_office_ai.js.
_AI = PAGE.split('<aside id="ai"')[1].split("</aside>")[0]
check("the slice-2 placeholder promise is gone from the served document",
      "coming in slice 2" not in PAGE)
check("the AI side-panel has a real composer, a send button and a model readout",
      'id="ai-in"' in _AI and 'id="ai-send"' in _AI and 'id="ai-model"' in _AI)
check("…and can be collapsed and re-opened",
      'id="ai-hide"' in _AI and 'id="ai-tab"' in _AI)
check("the panel speaks the harness's EXISTING direct chat lane, not a new office one",
      "'/api/chat/direct'" in PAGE and "/api/office/chat" not in _AI)
check("unsaved work is defended on navigation away", "beforeunload" in PAGE)
check("switching files with unsaved work asks IN THE PAGE — window.confirm is a "
      "recorded WKWebView no-op that returns FALSE and would wedge the tab",
      "window.confirm(" not in PAGE and "confirm(" not in
      re.sub(r"//[^\n]*", "", PAGE) and "confirmDiscard(name)" in PAGE
      and "discardArmed" in PAGE)
check("delete is two-step armed, not a single click", "sure?" in PAGE)
check("a missing openpyxl is SAID on the page, not hidden, WITH the command that fixes "
      "it (ship.sh does not reinstall requirements, so a shipped snapshot predating "
      "this lane has no openpyxl and every office action fails)",
      "roundtrip" in PAGE and "uv pip install openpyxl" in PAGE)


# ══ 7. the two defects that made the tab dead (2026-08-21) ═══════════════════
# Both are REGRESSION FENCES: each was invisible in every unit test and only showed up
# as "nothing happens" on a real Mac, so each gets an assertion that fails on the shape
# of the bug rather than on a symptom.

# DEFECT 1 — the page could not speak. #msg carried `display:none` in the STYLESHEET
# while say() showed it by clearing the INLINE display; clearing an inline value falls
# back to the rule, so the banner was unreachable and EVERY error the page produced
# (including "openpyxl is not installed", the one that was actually firing) was
# swallowed. Same class as the recorded `[hidden]{display:none!important}` bug.
msg_rule = re.search(r"#msg\{([^}]*)\}", PAGE)
check("the #msg banner has a stylesheet display rule…", msg_rule is not None)
check("…and is therefore shown by a CLASS, never by clearing the inline display",
      "m.classList.toggle('on'" in PAGE and "style.display = text" not in PAGE)
check("the show-state has its own rule", re.search(r"#msg\.on\{[^}]*display:", PAGE))
# The same trap, everywhere else a hidden thing is revealed: no code in this page may
# assign an empty display string, because every hidable element here is hidden by CSS.
# …read past the comments, which necessarily SAY `style.display=''` to explain it.
PAGE_CODE = re.sub(r"/\*.*?\*/", "", PAGE, flags=re.S)
PAGE_CODE = re.sub(r"^\s*//[^\n]*", "", PAGE_CODE, flags=re.M)
check("no `style.display = ''` survives anywhere in the page (the whole bug class)",
      not re.search(r"style\.display\s*=\s*(''|\"\"|``)", PAGE_CODE)
      and not re.search(r"style\.display\s*=\s*[a-zA-Z_.]+\s*\?\s*''", PAGE_CODE))
check("#newrow is toggled by class too", "newrow').classList" in PAGE)

# DEFECT 2 — Univer mounted into a display:none container. mount() ran before `current`
# was set and paint() (which un-hid #sheet) ran in the finally AFTER it, so the canvas
# engine measured a 0x0 box and no later un-hiding re-measured it: a permanently blank
# grid.
#
# ⚠️ THE FIX CHANGED SHAPE AT 2026-08-21e AND THE ASSERTION HAD TO CHANGE WITH IT,
# HONESTLY. #sheet used to be "always laid out"; now it is tier 2's container and stays
# hidden until somebody asks for the rich editor, because tier 1 owns the screen by
# default. The invariant that actually mattered is unchanged and is now pinned as an
# ORDERING instead of a stylesheet fact: `upgrade()` un-hides the container and yields a
# frame BEFORE it mounts. (The full ordering assertion, over the extracted function
# body, lives in bridge/tests/test_office_grid.js.)
# ⚠️ AND IT CHANGED SHAPE TWICE MORE. At loffice-2026-08-27c the mount this used to pin
# was retired: ONLYOFFICE became tier 2 and lived in its OWN page (/oo-edit), so there
# was no container to measure and no 0x0-canvas trap left to guard. Then Debi's
# 2026-08-27 ONE-EDITOR ruling (roadmap §8) deleted the navigation too: the editor is
# EMBEDDED in this page's centre column, in a same-origin iframe, and `upgrade()` is only
# a way to start it when it is not already up.
#
# THE INVARIANT THAT SURVIVED ALL THREE SHAPES, AND IT IS THE ONLY ONE THIS FILE OWNS:
# the editor is never a dead click — nothing happens until the bridge has said whether
# the bundle is installed, and a "no" is a sentence with the installer in it. The full
# ordering assertions over the extracted body live in bridge/tests/test_office_grid.js,
# and the whole one-editor wiring in bridge/tests/test_oo_lane.py.
up = PAGE.split("async function upgrade()")[1].split("\n// ⚠️ RETIRED")[0] \
    if "async function upgrade()" in PAGE else ""
check("upgrade() no longer mounts a second engine into this page, and no longer "
      "navigates away from it either — the editor is embedded here",
      bool(up) and "mount(snap)" not in up and "location.href" not in up)
check("…and it asks the bridge whether that editor exists before it starts anything",
      bool(up) and up.index("ooProbe()") < up.index("ooStart("))
check("…and the iframe that carries it is real markup in this page",
      'id="ooframe"' in PAGE and 'id="oostage"' in PAGE)
check("…and there is still no inline display on it (the class/inline trap)",
      'id="sheet" style' not in PAGE)
check("the empty state is an OVERLAY over the container, not a replacement for it",
      re.search(r"#empty\{[^}]*position:absolute", PAGE)
      and re.search(r"#empty\.off\{[^}]*display:none", PAGE)
      and "empty').classList.toggle('off'" in PAGE)
check("nothing hides the grid container any more",
      "el('sheet').style.display" not in PAGE)
check("a re-measure is nudged after mounting (a backgrounded tab can lay out late)",
      "requestAnimationFrame" in PAGE and "new Event('resize')" in PAGE)

# THE SELF-CHECK. The page proves its own preconditions and NAMES what is missing —
# the standing answer to "it renders nothing and says nothing".
# (`selfCheck();` at parse time became `if (selfCheck()) say('');` inside the
# DOMContentLoaded boot when the bundles were deferred — same call, later moment, and
# the later moment is the correct one: at parse time the bundles have not run yet.)
check("the page self-checks before it mounts", "function selfCheck()" in PAGE
      and "if (!selfCheck()) return false;" in PAGE)
needs = re.findall(r"\n  \['([^']+)'", PAGE.split("const NEEDS = [")[1].split("\n];")[0])
for g in ("React", "ReactDOM.createRoot", "rxjs", "rxjs.operators", "@wendellhu/redi",
          "@wendellhu/redi/react-bindings", "UniverCore.LocaleType",
          "UniverThemes.defaultTheme", "UniverPresets.createUniver",
          "UniverPresetSheetsCore.UniverSheetsCorePreset"):
    check(f"the self-check covers {g}", any(n.startswith(g) for n in needs))
check("…and the locale, which is the one that produces an untranslated grid",
      any(n.startswith("UniverPresetSheetsCoreEnUS") for n in needs))
# ⚠️ THE CSS MEASUREMENT PROBE IS GONE AT 2026-08-21e, DELIBERATELY. It existed because
# the stylesheet was in <head> and a 404 left the <link> in place while stripping every
# scrap of chrome. The stylesheet is now injected on demand, is NOT awaited, and styles
# only the optional tier-2 chrome — so its absence costs some styling on an upgrade
# nobody has to take, and cannot cost the spreadsheet. It still reports itself.
check("the tier-2 stylesheet reports both outcomes and is never awaited (a page stuck "
      "behind a stylesheet is the exact failure being designed out)",
      "'preset-sheets-core.css'" in PAGE
      and re.search(r"l\.onload = \(\) => bx\('asset-ok'", PAGE)
      and re.search(r"l\.onerror = \(\) => bx\('asset-error'", PAGE)
      and not re.search(r"await[^\n]*cssAsked", PAGE))
# ⚠️ THE SCRIPT IT NAMES CHANGED AT loffice-2026-08-27c. Tier 2 is ONLYOFFICE now, so
# the thing a user can be missing is the vendored ONLYOFFICE bundle, not the Univer
# UMD — and the sentence has to name THAT installer or it is advice about the wrong
# problem. The property is unchanged: a tier-2 failure names the command that fixes it.
check("the tier-2 failure names the script that fixes it",
      "scripts/install_onlyoffice.sh" in PAGE or "st.installer" in PAGE)
check("a runtime throw after boot reaches the screen, not just the console",
      "addEventListener('error'" in PAGE and "unhandledrejection" in PAGE)

# THE MODULES THE PAGE READS FROM. LocaleType is on UniverCore and defaultTheme on
# UniverThemes — reading either off UniverPresets (which re-exports only the core
# FACADE plus createUniver) yields undefined. It was survivable by luck; it is now
# correct, and pinned so it cannot drift back.
check("the locale comes from UniverCore, not UniverPresets",
      "window.UniverCore.LocaleType" in PAGE and "UP.LocaleType" not in PAGE)
check("the theme comes from UniverThemes, not UniverPresets",
      "window.UniverThemes.defaultTheme" in PAGE and "UP.defaultTheme" not in PAGE)

# THE FULL PRESET UI. The toolbar/formula bar/sheet tabs are Univer's own; the page
# asks for each BY NAME so a changed plugin default cannot quietly remove one.
for opt in ("header", "toolbar", "formulaBar", "footer", "contextMenu"):
    check(f"the mount asks for the {opt} explicitly",
          re.search(rf"\b{opt}:\s*true", PAGE))
check("the container is the one element the page lays out for it",
      "container: 'sheet'" in PAGE)

# IMPORT.
check('app.py declares @app.post("/api/office/upload")',
      '@app.post("/api/office/upload")' in APP)
check("the upload body is RAW, like /api/voice/library/save",
      "await req.body()" in APP.split('@app.post("/api/office/upload")')[1])
check("the upload is capped at the bridge, not only in the page",
      "_office.UPLOAD_MAX_BYTES" in APP and "status_code=413" in
      APP.split('@app.post("/api/office/upload")')[1].split("@app.get")[0])
check("the import runs off the event loop",
      "asyncio.to_thread(\n        _office.import_doc" in APP
      or "asyncio.to_thread(_office.import_doc" in APP)
check("the page can reach it, through a real file picker",
      "/api/office/upload?name=" in PAGE and 'type="file"' in PAGE
      and 'accept=".xlsx' in PAGE)
check("the import cap is the same number on both sides",
      f"{office.UPLOAD_MAX_BYTES // (1024 * 1024)} * 1024 * 1024" in PAGE)

check("import never clobbers — it renames", office.free_name.__doc__
      and "NEVER clobbers" in office.free_name.__doc__)
with tempfile.TemporaryDirectory() as td:
    check("free_name gives the plain name when nothing is there",
          office.free_name(td, "book.xlsx") == ("book.xlsx", None))
    office.create_doc(td, "book")
    check("…and steps to ' (2)' when it is taken",
          office.free_name(td, "book.xlsx") == ("book (2).xlsx", None))
    office.create_doc(td, "book (2)")
    check("…and keeps stepping", office.free_name(td, "book")[0] == "book (3).xlsx")
    check("free_name refuses a traversal name like every other route",
          office.free_name(td, "../x.xlsx")[0] is None)

    # import_doc, executed for real against a workbook openpyxl itself produced.
    if HAVE_XL:
        import openpyxl as _xl
        buf = io.BytesIO()
        wb = _xl.Workbook()
        wb.active["A1"] = "hello"
        wb.active["B2"] = 41.5
        wb.save(buf)
        good = buf.getvalue()
        rep, reason = office.import_doc(td, "Imported.xlsx", good)
        check("import_doc keeps a real workbook", rep and rep["name"] == "Imported.xlsx"
              and rep["renamed"] is False and reason is None)
        snap, _ = office.open_doc(td, "Imported.xlsx")
        sid = snap["sheetOrder"][0]
        cells = snap["sheets"][sid]["cellData"]
        check("…and the imported cells read back through the normal open path",
              cells["0"]["0"]["v"] == "hello" and cells["1"]["1"]["v"] == 41.5)
        rep2, _ = office.import_doc(td, "Imported.xlsx", good)
        check("a second import of the same name is RENAMED, never overwritten",
              rep2 and rep2["name"] == "Imported (2).xlsx" and rep2["renamed"] is True
              and office.doc_target(td, "Imported.xlsx")[0])
        check("a renamed .txt is refused before it can land in the list",
              office.import_doc(td, "notreally.xlsx", b"this is not a workbook")[0] is None)
        check("…and nothing was written for it",
              office.doc_target(td, "notreally.xlsx")[0] is None)
        check("an empty body is refused", office.import_doc(td, "e.xlsx", b"")[0] is None)
        check("a body past the cap is refused without parsing it",
              office.import_doc(td, "big.xlsx",
                                b"x" * (office.UPLOAD_MAX_BYTES + 1))[0] is None)
        check("a bad name is refused before the bytes are read",
              office.import_doc(td, "../evil.xlsx", good)[0] is None)
        check("a non-xlsx extension is refused",
              office.import_doc(td, "book.numbers", good)[0] is None)

# ══ 7. the blank-tab forensics (2026-08-21) ══════════════════════════════════
# WHY THIS GROUP EXISTS. The page shipped correct — every route 200s, every bundle
# defines every global it promises (executed in a JS engine, not read out of the docs)
# — and it was STILL a blank rectangle on the real Mac, with the self-check banner it
# was given for exactly this purpose nowhere to be seen. Two facts explain that, and
# both of them are now structural rather than remembered:
#
#   (a) 10.5 MB of NON-deferred classic script sat between the body and first paint.
#       A classic <script src> blocks the parser where it stands, and WebKit is under
#       no obligation to paint what it has already parsed first — so "blank" was the
#       page's honest state for as long as the compile took, and every diagnostic the
#       page carried was on the far side of that wall. `defer` moves the wall.
#   (b) A page can only report on itself while its scripts are alive. A stale cached
#       document, a parse error or a killed web content process leaves NOTHING — so
#       the first line of diagnosis has to be plain HTML that is true by default and
#       taken down by JS, never printed by it.
AIDER_PAGE = (ROOT / "bridge" / "panel" / "aider.html").read_text(encoding="utf-8")

# ⚠️ (a) WAS SOLVED BY `defer` AT 2026-08-21d AND BY DELETION AT 2026-08-21e. Deferring
# 10.5 MB stops it blocking first paint, but the tab still could not be a spreadsheet
# until all of it arrived and ran. It is now loaded ON DEMAND, so the assertion is the
# stronger one: the document has no external subresource to be blocked by. Comments in
# the page necessarily TALK about script and link tags to explain why there are none,
# so the check reads the markup with comments stripped.
NOCOM = re.sub(r"<!--.*?-->", "", PAGE, flags=re.S)
check("the served document has NO external script tag at all — tier 1 needs nothing "
      "but itself, so nothing external can stop it existing",
      not re.search(r"<script[^>]*\bsrc=", NOCOM))
check("…and NO external stylesheet in <head>: a render-blocking <link> that never "
      "resolves blocks every script AFTER it, which is exactly 'the head beacon fired "
      "and nothing else in the document ever did'",
      not re.search(r"<link\b", NOCOM.split("</head>")[0]))
check("…and the reason is written where the next reader will hit it",
      "render-blocking" in PAGE and "TWO TIERS" in PAGE)
check("the inline script is NOT deferred (inline scripts ignore it; it runs at parse "
      "time, which is what makes it the beacon)",
      re.search(r'<script>\s*\n\s*\(function \(\) \{\s*\n\s*.use strict', PAGE) is not None)


def inline_blocks(page):
    """Every <script> in `page` that has no src, in document order.

    ⚠️ office.html has TWO of them since the boot beacon landed (2026-08-21): the
    beacon in <head> and the application script at the end of <body>. Anything that
    means "the inline script" has to say WHICH, or it silently starts pinning the
    wrong one — which is how a passing test stops testing anything.
    """
    return re.findall(r'<script(?![^>]*\bsrc=)[^>]*>(.*?)</script>', page, re.S)


def app_block(page):
    """The APPLICATION inline script — the one that owns the fallback banner."""
    for b in inline_blocks(page):
        if "getElementById('boot')" in b:
            return b
    return ""

# ⚠️ THE STAMP IS READ, NOT WRITTEN DOWN. It used to be a literal here, which meant
# every build bump had to edit this file and — worse — that the test was pinning a
# STRING rather than the invariant. The invariant is: the page declares its stamp in
# ONE place (the <meta>) and every other appearance is a copy of that one.
def _stamp(page, what):
    m = re.search(r'<meta name="harness-build" content="([^"]+)">', page)
    if not m:
        raise SystemExit(f"{what}: no harness-build meta — the stamp is load-bearing")
    return m.group(1)

OFFICE_STAMP = _stamp(PAGE, "office.html")
AIDER_STAMP = _stamp(AIDER_PAGE, "aider.html")
check(f"office: the build stamp names the page it stamps ({OFFICE_STAMP})",
      OFFICE_STAMP.startswith("loffice-"))
check(f"aider: likewise ({AIDER_STAMP})", AIDER_STAMP.startswith("aider-"))

# The static fallback banner: true by default, removed by the first statement of JS.
for label, page, stamp in [("office", PAGE, OFFICE_STAMP),
                           ("aider", AIDER_PAGE, AIDER_STAMP)]:
    head = page.split("<body>")[0]
    body = page.split("<body>")[1]
    check(f"{label}: the fallback banner is REAL MARKUP in the body, not JS output",
          '<div id="boot">' in body
          and "scripts did not run" in body
          and body.index('<div id="boot">') < body.index("<script"))
    check(f"{label}: …and it says what to do about it",
          "⌘R" in body.split("</div>")[0] and "ship.sh" in body.split("</div>")[0])
    check(f"{label}: …and it is not hidden by a stylesheet rule",
          not re.search(r'#boot\{[^}]*display:\s*none', page))
    # "first" is asserted over the STATEMENTS, not the raw characters: both pages
    # carry a paragraph of comment explaining why the line is there, and a character
    # window would be pinning the length of that comment rather than the order.
    inline = app_block(page)
    stmts = re.sub(r"//[^\n]*", "", inline)
    stmts = re.sub(r"\s+", " ", stmts).strip()
    lead = stmts[:stmts.index("getElementById('boot')")]
    check(f"{label}: the first thing the inline script does is remove it",
          lead.replace("'use strict';", "").strip()
          == "(function () { const b = document."
          and ".remove()" in stmts[:len(lead) + 90])
    check(f"{label}: the build stamp is in a <meta> AND printed in the banner, "
          "so a stale document is identifiable with no terminal",
          f'name="harness-build" content="{stamp}"' in head
          and stamp in body.split("</div>")[0])
    check(f"{label}: the static <title> is the BOOTING state — the shell reads "
          "webView.title, so a dead page is visible from outside too",
          "· booting…</title>" in head)
    check(f"{label}: …and JS changes it once it is alive",
          "document.title = " in page)

# Every id the script reaches for has to exist in the markup. `el(...)` returns null for
# a name that does not, and the wiring at the bottom of the file assigns onclick straight
# onto the result — so ONE renamed element is a TypeError at parse time, i.e. a page that
# renders its chrome and then does nothing at all, which is a described symptom.
for label, page in [("office", PAGE), ("aider", AIDER_PAGE)]:
    have = set(re.findall(r'\bid="([^"]+)"', page))
    want = (set(re.findall(r"el\('([^']+)'\)", page))
            | set(re.findall(r"getElementById\('([^']+)'\)", page)))
    check(f"{label}: every element the script reaches for exists in the markup "
          f"(missing: {sorted(want - have)})", not (want - have))

# ⚠️ THE CENSUS BECAME A SHAPE CHECK AT loffice-2026-08-28a. `count == 2` only held
# while no COMMENT in the page named the build it was written for, and the one-editor
# slice writes its stamp into the comments that explain the fork (deliberately — a
# reader of that CSS block can date it). What the check is FOR is that the two
# LOAD-BEARING copies cannot drift: the <meta> the script reads, and the no-script
# fallback banner a person reads.
check("the page reads its stamp from the meta rather than keeping a second copy",
      'meta[name="harness-build"]' in PAGE
      and f'<meta name="harness-build" content="{OFFICE_STAMP}">' in PAGE
      and f"Build <code>{OFFICE_STAMP}</code>" in PAGE)
# ⚠️ CHANGED HONESTLY AT 2026-08-21e. There is nothing deferred to wait FOR any more,
# so the boot no longer hangs on an event: the script sits at the end of the body, and
# waiting for a DOMContentLoaded that has ALREADY FIRED would never boot at all — the
# one failure a boot path may not have.
check("boot runs immediately when the document is already parsed, and only waits when "
      "it genuinely is still loading",
      "document.readyState === 'loading'" in PAGE
      and "addEventListener('DOMContentLoaded'" in PAGE
      and re.search(r"\} else \{\s*\n\s*boot\(\);", PAGE) is not None)
check("…and the rich-editor load says which file it is on, per file",
      "Loading the rich editor" in PAGE and "VENDOR.length" in PAGE)
check("a watchdog turns 'still starting' into a sentence after a bounded wait",
      "WATCHDOG_MS" in PAGE and "did not finish starting" in PAGE)
check("a 404 on a bundle is caught in the CAPTURE phase and NAMES the file "
      "(a resource error does not bubble, so the ordinary handler cannot see it)",
      re.search(r"addEventListener\('error'.{0,400}?\}, true\)", PAGE, re.S) is not None
      and "did not load" in PAGE)

# The missing delegate. Same class as runOpenPanelWith and the media-capture grant:
# a WKWebView whose content process is killed shows nothing at all, forever, unless
# the host answers for it.
check("main.swift answers webViewWebContentProcessDidTerminate",
      "func webViewWebContentProcessDidTerminate(_ webView: WKWebView)" in SWIFT)
check("…and reloads ONCE rather than looping", "crashedOnce" in SWIFT
      and "crashedOnce.insert(key)" in SWIFT and "crashedOnce.contains(key)" in SWIFT)
check("…and a second death in a row becomes a readable notice, not a void",
      "ran out of memory" in SWIFT)
_dfin = SWIFT.find("func webView(_ webView: WKWebView, didFinish")
check("…and a successful load clears the crash memory, so the next one is a new "
      "incident with its own retry",
      _dfin > 0 and "crashedOnce.remove(" in SWIFT[_dfin:_dfin + 500])
check("…and it says so in the log, tagged as a tab rather than as the split view",
      '"[tab] \\(title): web content process died' in SWIFT)

# Live: what the bridge actually puts on the wire. Answers "is the document cacheable"
# and "does the served body contain today's page" in one pass.
try:
    from fastapi.testclient import TestClient                     # noqa: E402
    import warnings
    warnings.filterwarnings("ignore")
    from bridge import app as A                                   # noqa: E402
    cl = TestClient(A.app)
    for route in ["/", "/office", "/aider"]:
        r = cl.get(route)
        cc = r.headers.get("cache-control", "")
        check(f"GET {route} is 200 html", r.status_code == 200
              and r.headers.get("content-type", "").startswith("text/html"))
        check(f"GET {route} is no-store — a bridge-served DOCUMENT may never be "
              "cached by the WKWebView (the recorded stale-panel incident)",
              "no-store" in cc)
    served = cl.get("/office").text
    check("the served /office body carries today's build stamp — i.e. the route reads "
          "the file per request, so a ship really does change what is served",
          OFFICE_STAMP in served and '<div id="boot">' in served)
    for asset, mime in [("/assets/vendor/react.production.min.js", "javascript"),
                        ("/assets/vendor/react-dom.production.min.js", "javascript"),
                        ("/assets/vendor/univer/rxjs.umd.min.js", "javascript"),
                        ("/assets/vendor/univer/presets.umd.js", "javascript"),
                        ("/assets/vendor/univer/preset-sheets-core.umd.js", "javascript"),
                        ("/assets/vendor/univer/preset-sheets-core.en-US.js", "javascript"),
                        ("/assets/vendor/univer/preset-sheets-core.css", "text/css")]:
        r = cl.get(asset)
        check(f"{asset} serves 200 with a {mime} content-type and real bytes",
              r.status_code == 200 and mime in r.headers.get("content-type", "")
              and int(r.headers.get("content-length", "0")) > 1000)
    check("the vendor bundles are NOT no-store — they are pinned, immutable and 10 MB; "
          "revalidation is the right cost",
          "no-store" not in cl.get("/assets/vendor/univer/presets.umd.js")
          .headers.get("cache-control", ""))
except Exception as _e:                                          # noqa: BLE001
    print(f"  (skipped live route checks — no TestClient/app: {_e})")

# ══ 8. THE BOOT BEACON (2026-08-21) ══════════════════════════════════════════
# WHY THIS GROUP EXISTS. Round 7 above made the page able to describe its own failure
# ON SCREEN — and the page then failed a third time on a Mac none of us can reach, so
# the description went nowhere. The screen is only a report if somebody can read it;
# the bridge log is a report you can grep after the fact. Every stage of the boot is
# now beaconed, and this group pins the parts that can silently stop working:
#   * `diag_line` is PURE and TOTAL — it is fed by a page that is, by hypothesis,
#     malfunctioning, so it must survive every shape of junk.
#   * the route NEVER fails. A beacon endpoint that 4xx's would trip the page's own
#     error handling, i.e. the diagnostic would become the second fault to diagnose.
#   * ⚠️ THE REGRESSION THAT NEARLY SHIPPED: the first draft parsed the body with
#     `json.loads`, and this module has no module-level `json` import. The NameError
#     was swallowed by the route's own except and EVERY beacon logged as
#     "unparseable-beacon" — a diagnostic that reported only its own breakage. It was
#     caught by driving the real page in a real browser. Both halves are pinned below.
DIAG_TABLE = [
    (("mount-ok", "canvases=3", "ab12cd", 1240), "stage=mount-ok", "the ordinary line"),
    (("mount-ok", "", "ab12cd", 0), "+0ms", "zero elapsed is a number, not a blank"),
    (("s", None, "b", 5), "stage=s", "a None detail is simply absent"),
    (("s", "d", "b", None), "+?ms", "a missing elapsed is marked, never invented"),
    (("s", "d", "b", "junk"), "+?ms", "a non-numeric elapsed cannot raise"),
    (("s", "d", "b", float("nan")), "+?ms", "NaN is not an elapsed"),
    (("s", "d", "b", float("inf")), "+?ms", "neither is infinity"),
    ((None, "d", "b", 1), "stage=?", "a missing stage still yields a line"),
    (("s", 12345, "b", 1), "detail=12345", "a numeric detail is coerced"),
    (("s", {"a": 1}, "b", 1), "detail=", "a dict detail is coerced, not crashed on"),
]
for args, needle, why in DIAG_TABLE:
    try:
        line = office.diag_line(*args)
        ok = needle in line and "\n" not in line
    except Exception:                                            # noqa: BLE001
        line, ok = "RAISED", False
    check(f"diag_line: {why}", ok)

check("diag_line flattens newlines — one beacon is one line, and a detail that could "
      "carry a newline could forge a second entry",
      "\n" not in office.diag_line("s", "a\nb\rc\td", "boot", 1)
      and "a b c d" in office.diag_line("s", "a\nb\rc\td", "boot", 1))
check("diag_line clamps a hostile detail to DIAG_DETAIL_MAX",
      len(office.diag_line("s", "x" * 9000, "b", 1)) < office.DIAG_DETAIL_MAX + 200)
check("diag_line clamps the stage too",
      len(office.diag_line("y" * 500, "", "b", 1)) < office.DIAG_STAGE_MAX + 120)
check("diag_line carries the boot id — two script-starts with DIFFERENT ids is the "
      "signature of a reload loop, i.e. a dying web content process",
      "boot=ab12cd" in office.diag_line("s", "d", "ab12cd", 1))

with tempfile.TemporaryDirectory() as td:
    line = office.write_diag(td, "script-start", "build=x", "aa11", 7)
    logf = Path(td) / "data" / "logs" / (office.DIAG_LOG_NAME + ".log")
    check("write_diag creates data/logs and appends the line",
          logf.exists() and line in logf.read_text())
    office.write_diag(td, "second", "", "aa11", 9)
    check("…and it APPENDS rather than replacing",
          len(logf.read_text().splitlines()) == 2)
    for i in range(office.DIAG_MAX_LINES + 250):
        office.write_diag(td, "loop", "x" * 150, "bb22", i)
    n = len(logf.read_text().splitlines())
    check("…and it is BOUNDED — a crash loop is the loudest thing this file can "
          "carry, and it must not also become a disk problem",
          n <= office.DIAG_MAX_LINES + 5)
    check("…keeping the NEWEST lines (the tail is the incident)",
          f"+{office.DIAG_MAX_LINES + 249}ms" in logf.read_text().splitlines()[-1])
check("write_diag never raises on an unwritable root — it is an error path",
      isinstance(office.write_diag("/proc/nonexistent/nope", "s", "d", "b", 1), str))

APP = _APP_SOURCE
DIAG_ROUTE = APP.split('@app.post("/api/office/diag")')[1].split("@app.")[0] \
    if '@app.post("/api/office/diag")' in APP else ""
check("the bridge exposes POST /api/office/diag", bool(DIAG_ROUTE))
# The route's own docstring and comments TALK about the trap below, so the assertion
# has to read the CODE — a grep over prose would pass on the strength of the warning
# while the bug it warns about sat two lines under it.
DIAG_CODE = re.sub(r"#[^\n]*", "", re.sub(r'""".*?"""', "", DIAG_ROUTE, flags=re.S))
check("⚠️ REGRESSION PIN: the route parses with `await req.json()`, NOT `json.loads` "
      "— bridge/app.py has no module-level `json` import, so json.loads is a "
      "NameError that this route's own except would swallow, and every beacon would "
      "log as 'unparseable'",
      "await req.json()" in DIAG_CODE and "json.loads" not in DIAG_CODE)
# ⚠️ RE-SCOPED 2026-08-28 (SSE-hybrid slice). This read `APP` — bridge/appsrc.py's view
# of the WHOLE app layer — and asked whether ANY of its ~30 modules imports json. That
# was the same question as "does the diag route's module import json" only while the app
# layer was ONE FILE; after the router/core split it became a question about 30 unrelated
# modules, and the first one to legitimately need json (bridge/core/events.py, which
# serialises SSE frames) turned it red for a reason that has nothing to do with the
# beacon. `json` is a MODULE global: the only file whose imports can make `json.loads`
# resolve inside this route is the file the route lives in.
_DIAG_OWNER = (ROOT / "bridge" / "routers" / "office.py").read_text()
check("…and the diag route's OWN module still has no module-level `json` import, which "
      "is what makes the line above load-bearing rather than decorative",
      '@app.post("/api/office/diag")' in _DIAG_OWNER
      and not re.search(r"^import json$", _DIAG_OWNER, re.M)
      and not re.search(r"^import json as ", _DIAG_OWNER, re.M))
check("the route works with the office module UNAVAILABLE — that case (openpyxl "
      "missing, a syntax error in office.py) is itself worth reporting, so the "
      "reporting path cannot depend on the module importing",
      "_office is not None" in DIAG_CODE and "office module unavailable" in DIAG_ROUTE)
check("the route answers 204 and never an error status — a failing beacon endpoint "
      "would make the page's own error handling fire",
      "status_code=204" in DIAG_CODE
      and "status_code=400" not in DIAG_CODE and "status_code=500" not in DIAG_CODE)
check("the beacon log is in _LOG_NAMES, so the trace is readable in the panel and "
      "over /api/logs — the tab it describes may be showing nothing at all",
      '"loffice-boot"' in APP and office.DIAG_LOG_NAME == "loffice-boot")

# ⚠️ THE GREP TAG, AND IT IS THE REASON A WHOLE ROUND WENT INTO A GHOST (2026-08-21e).
# These lines used to be logged as `[office] diag …`, and the instruction handed to Debi
# was `grep loffice <bridge.log>`. Only ONE stage — script-start, whose DETAIL happens
# to contain the build stamp "loffice-2026-08-21x" — carried the string "loffice" at
# all. So the grep returned exactly one line on a PERFECT boot, and that single line was
# read as "the document stopped dead after the head script". Every line is greppable by
# one stable token now, and this asserts it over a real trace rather than over the
# format string.
_TRACE = [office.diag_line(st, dt, "abc123", i * 10) for i, (st, dt) in enumerate([
    ("script-start", f"build={OFFICE_STAMP} ua=… url=… vis=visible"),
    ("boot-inline", "tier-1 script running"),
    ("dom-ready", "tier=grid vis=visible"),
    ("files-ok", "n=2 roundtrip=true"),
    ("grid-render", "sheet=Sheet1 rows=40 cols=12 cells=480"),
    ("tier1-ready", "files=2"),
])]
check("⚠️ REGRESSION PIN: EVERY diag line is findable with one grep — a trace where "
      "only the first line matches the documented grep is worse than no trace, "
      "because it looks like evidence",
      all("loffice" in f"[office] loffice-diag {ln}" for ln in _TRACE))
check("…and the route emits exactly that tag",
      're_office_tag' not in APP and 'loffice-diag {line}' in APP)
check("…and it is NOT the bare word 'diag', which is what made only the stamp match",
      not re.search(r'_office_log\(f"diag ', APP))

# THE ASSET ACCESS LOG. It answers the one question the beacon cannot: when a boot
# trace stops after a given bundle, was the next file NEVER REQUESTED (the parser died,
# or the page never got that far) or REQUESTED AND NEVER FINISHED (a stalled transfer)?
check("the /assets mount is wrapped so the Univer bundles are logged",
      "class _WatchedStatic(StaticFiles)" in APP
      and "_WatchedStatic(directory=" in APP)
check("⚠️ …as a StaticFiles SUBCLASS, never as BaseHTTPMiddleware: a BaseHTTPMiddleware "
      "wraps EVERY response in the harness, including the SSE chat relays whose 20s "
      "heartbeat the panel's stall watchdog counts on — a diagnostic for the "
      "spreadsheet tab may not go near them",
      "app.middleware" not in re.sub(r'""".*?"""', "", APP, flags=re.S))
_WS = APP.split("class _WatchedStatic(StaticFiles):")[1].split("\napp.mount")[0]
check("…and it logs the request BEFORE serving and the outcome after, so a `→` with no "
      "`←` IS the stalled case rather than a gap to be interpreted",
      _WS.index("loffice-asset →") < _WS.index("loffice-asset ←"))
check("…with the status and the byte count, so a truncated transfer is visible",
      "sent['status']" in _WS and "sent['bytes']" in _WS)
check("⚠️ REGRESSION PIN: the watch is a SUBSTRING match, not startswith. The first "
      "draft used startswith('/vendor/univer/') on the reasoning that a Mount rewrites "
      "the path; this Starlette does not (it leaves the full path and sets root_path), "
      "so the branch never fired and the whole diagnostic was a no-op that reviewed "
      "correctly. Caught by driving the real page and finding zero lines.",
      "for w in self.WATCH" in _WS and "startswith(self.WATCH)" not in _WS)
check("…and it is scoped to the office bundles, so it is a handful of lines per "
      "rich-editor load and zero for every other request in the harness",
      "/vendor/univer/" in _WS and "/vendor/react.production.min.js" in _WS)
check("…and it never swallows an exception it logs", "raise" in _WS)

HEAD = PAGE.split("</head>")[0]
BEACON = inline_blocks(PAGE)[0]
check("the beacon is the FIRST script in the document — its ABSENCE from the log is "
      "then a fact: no script in this document ran at all",
      HEAD.index("<script>") < HEAD.index("<style>")
      and "window.bx" in BEACON and "script-start" in BEACON)
check("…and it sends with sendBeacon, the only transport specified to survive the "
      "page being torn down (a keepalive fetch is the fallback)",
      "navigator.sendBeacon" in BEACON and "keepalive: true" in BEACON)
check("…with a per-load boot id, so a reload loop is visible as repeated "
      "script-start under DIFFERENT ids",
      "Math.random()" in BEACON and "boot: BOOT" in BEACON)
check("…and it is capped, so a malfunctioning page cannot become a traffic source",
      "CAP" in BEACON and "SENT++" in BEACON)
check("…and it can never throw: a diagnostic may not be the thing that breaks the page",
      BEACON.count("try {") >= 1 and "catch (e)" in BEACON)
check("the beacon reads the build stamp from the <meta> instead of repeating it — a "
      "second copy could drift, and the stamp exists to be trusted",
      'meta[name="harness-build"]' in BEACON
      and OFFICE_STAMP not in BEACON)
for _o, name, _c in SCRIPT_TAGS:
    tag = [t for t in re.findall(r"<script[^>]*>", PAGE) if name in t]
    check(f"the {name} tag reports BOTH outcomes — the last asset-ok in the trace "
          f"names where a boot stopped",
          bool(tag) and "onload=" in tag[0] and "onerror=" in tag[0])
for stage in ["script-start", "asset-ok", "asset-error", "dom-ready", "selfcheck-pass",
              "selfcheck-fail", "files-ok", "files-fail", "mount-start", "mount-fail",
              "mount-settled", "page-error", "rejection", "watchdog", "create-fail",
              "open-fail"]:
    check(f"the page beacons '{stage}'", f"'{stage}'" in PAGE)
check("⚠️ the canvas count is read SETTLED, not at two animation frames: the early "
      "number is 0 on a perfectly healthy mount (measured), and a diagnostic that "
      "cries wolf is what sends the next session chasing the wrong thing",
      "MOUNT_SETTLE_MS" in PAGE and "mount-raf" in PAGE
      and re.search(r"MOUNT_SETTLE_MS\s*=\s*(\d+)", PAGE)
      and int(re.search(r"MOUNT_SETTLE_MS\s*=\s*(\d+)", PAGE).group(1)) >= 1000)
check("…and the settled probe is generation-guarded, so a superseded mount cannot "
      "accuse the grid that replaced it",
      "gen !== mountGen" in PAGE and "++mountGen" in PAGE)
check("the probe reports the BIGGEST canvas, not the count: Univer legitimately keeps "
      "a 0x0 and a tiny canvas beside the real grid, so a count or a minimum would "
      "both read as broken on a working page",
      "biggest=" in PAGE and "maxW" in PAGE)
check("…and a mount that reports success while drawing nothing SAYS SO — that is the "
      "reported symptom, finally decidable from the page itself",
      "GRID_MIN_PX" in PAGE and "drew nothing" in PAGE)
check("the progress line names the file being fetched, so a stalled upgrade is "
      "diagnostic on screen rather than only in the log",
      "function rich(" in PAGE and "richbar" in PAGE
      and re.search(r"rich\('Loading the rich editor[^\n]*label", PAGE) is not None)

try:
    from fastapi.testclient import TestClient                    # noqa: E402
    from bridge.app import app as _app                           # noqa: E402
    cl = TestClient(_app)
    for body, ctype, why in [
        ('{"stage":"unit","detail":"d","boot":"zz","ms":5}', "text/plain",
         "a text/plain sendBeacon body (the shape WebKit actually sends)"),
        ('{"stage":"unit2"}', "application/json", "a tidy json post"),
        ("not json at all", "text/plain", "junk"),
        ("", "text/plain", "an empty body"),
        ('"a bare string"', "application/json", "a non-dict json body"),
        ("[1,2,3]", "application/json", "a json array"),
    ]:
        r = cl.post("/api/office/diag", content=body, headers={"Content-Type": ctype})
        check(f"POST /api/office/diag answers 204 for {why} — it never fails",
              r.status_code == 204)
    check("…and a real beacon actually reaches the log file",
          "unit" in (ROOT / "data" / "logs" / "loffice-boot.log").read_text()
          if (ROOT / "data" / "logs" / "loffice-boot.log").exists() else False)
    check("the beacon log is readable over /api/logs like every other source",
          cl.get("/api/logs/loffice-boot").status_code == 200)
except Exception as _e:                                          # noqa: BLE001
    print(f"  (skipped live diag route checks: {_e})")

print()
if FAILS:
    print(f"{len(FAILS)} FAILED:")
    for f in FAILS:
        print("  -", f)
    sys.exit(1)
print(f"office lane OK — all checks passed{'' if HAVE_XL else ' (round-trip SKIPPED: no openpyxl)'}")

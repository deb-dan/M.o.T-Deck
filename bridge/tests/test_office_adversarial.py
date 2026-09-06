"""ADVERSARIAL QA CATALOGUE for the LOffice server lane — bridge/office.py and
bridge/office_ops.py.

⚠️ THIS FILE IS NOT A GATE AND MUST NEVER BE WIRED INTO ONE. It EXITS 0 ALWAYS.

WHAT IT IS: an executable catalogue of behaviours that are silently wrong from a user's
point of view, written 2026-08-28 alongside docs/research/
2026-08-28-adversarial-findings-server.md. Every finding in that document has a repro
here, in the same order, under the same F-nn id. A batch-fix builder runs this file, sees
which findings still REPRODUCE, fixes them, and runs it again to watch them flip to
FIXED. When every line says FIXED, delete the file (or promote the interesting halves
into bridge/tests/test_office_mcp.py, where they can be a gate).

    data/bridge-venv/bin/python bridge/tests/test_office_adversarial.py

Each check prints ONE line:

    FINDING F-03 [LIES-TO-USER]   REPRODUCES  one-line title
    FINDING F-03 [LIES-TO-USER]   fixed       one-line title

REPRODUCES means the defect is still there. `fixed` means the probe no longer sees it —
which for a few of these could also mean the probe drifted, so a fixer who sees an
unexpected `fixed` should read the finding's paragraph in the doc before believing it.

A handful of lines are marked [CONTROL] and print `holds` rather than `fixed`. Those are
the parts of the lane that are ALREADY RIGHT (the cap boundaries, the malformed-reference
refusals, all-or-nothing, read-side sheet targeting), pinned here so that a batch fix to
the 34 real findings cannot quietly regress them. A CONTROL that prints BROKEN is a new
bug the fix introduced.

THE ONE PROBE THAT CANNOT BE RUN INLINE is F-18 (an infinite loop in
office_ops.add_sheet). It is guarded by SIGALRM so this file cannot hang; on a platform
without SIGALRM it is reported as SKIPPED rather than run.
"""
from __future__ import annotations

import datetime as _dt
import json
import os
import shutil
import signal
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT))

from bridge import office, office_ops as oo   # noqa: E402

try:
    import openpyxl
except Exception as _e:                                          # noqa: BLE001
    print(f"SKIPPED ENTIRELY: openpyxl is not importable ({_e}) — run this with "
          "data/bridge-venv/bin/python")
    sys.exit(0)


# ══ MOT Deck ═══════════════════════════════════════════════════════════════
SEV = {}
SEEN = []


def finding(fid, sev, title, reproduces, detail=""):
    """Record one catalogue line. `reproduces` is TRUTHY when the defect is still there."""
    SEV.setdefault(sev, {"reproduces": 0, "fixed": 0, "skipped": 0})
    if reproduces == "skip":
        SEV[sev]["skipped"] += 1
        state = "SKIPPED   "
    elif sev == CTL:
        # A CONTROL: `reproduces` truthy means the behaviour BROKE. These are here so a
        # batch fix cannot quietly regress the parts of the lane that are already right.
        if reproduces:
            SEV[sev]["reproduces"] += 1
            state = "BROKEN    "
        else:
            SEV[sev]["fixed"] += 1
            state = "holds     "
        SEEN.append(fid)
        print(f"FINDING {fid} [{sev:16}] {state}  {title}")
        if reproduces:
            for line in str(detail).splitlines():
                print(f"                                       · {line}")
        return
    elif reproduces:
        SEV[sev]["reproduces"] += 1
        state = "REPRODUCES"
    else:
        SEV[sev]["fixed"] += 1
        state = "fixed     "
    SEEN.append(fid)
    print(f"FINDING {fid} [{sev:16}] {state}  {title}")
    if detail and reproduces and reproduces != "skip":
        for line in str(detail).splitlines():
            print(f"                                       · {line}")


LIE = "LIES-TO-USER"
REF = "REFUSES-WRONGLY"
COS = "COSMETIC"
CTL = "CONTROL"        # already-correct behaviour, pinned so a batch fix cannot regress it

R = tempfile.mkdtemp(prefix="loffice-advqa-")
OD = os.path.join(R, "data", "office")
os.makedirs(OD, exist_ok=True)


def wb_new(name, build=None):
    """A fresh workbook in the scratch office dir, built by openpyxl (i.e. by something
    that is NOT this lane — the point is to start from a real file)."""
    oo.changeset_clear()
    oo.heartbeat_clear()
    p = os.path.join(OD, name)
    if os.path.exists(p):
        os.remove(p)
    wb = openpyxl.Workbook()
    if build:
        build(wb)
    wb.save(p)
    return p


def stage(name, ops, sheet=None):
    return oo.stage_changes(R, "advqa", name, sheet, ops)


def apply(cid):
    """(out, reason). An UNCAUGHT exception is returned as a reason prefixed 'EXC ' —
    which is itself a finding: the route has no handler for it and answers 500."""
    try:
        return oo.apply_changeset(R, cid)
    except Exception as e:                                       # noqa: BLE001
        return None, f"EXC {type(e).__name__}: {e}"


def sa(name, ops, sheet=None):
    out, reason = stage(name, ops, sheet)
    if out is None:
        return None, None, reason
    a, r2 = apply(out["changeset_id"])
    return out, a, r2


def load(name, sheet=None):
    wb = openpyxl.load_workbook(os.path.join(OD, name))
    return wb[sheet] if sheet else wb.active


print(__doc__.splitlines()[0])
print("=" * 100)


# ══ 1. VALUE TYPING ═══════════════════════════════════════════════════════════
print("\n── 1. value typing ──")

# F-09 — an explicit TEXT number format (@) does not stop numeric coercion, and a custom
# pattern is replaced rather than kept.
def _b(wb):
    ws = wb.active
    ws.cell(1, 1).number_format = "@"                 # Debi: "this column is TEXT"
    ws.cell(2, 1, 5)
    ws.cell(2, 1).number_format = '0.00" kg"'         # Debi: "these are kilos"
wb_new("f09.xlsx", _b)
out, a, _r = sa("f09.xlsx", [{"op": "set", "at": "A1",
                              "values": [["1,200"], ["1,200"]]}])
ws = load("f09.xlsx")
# FIXED means BOTH halves of the finding's "correct behaviour": the `@` cell is not
# coerced at all, AND a replaced pattern is named on the card.
_f09_notes = " ".join((out or {}).get("notes") or [])
finding("F-09", LIE,
        "numeric coercion overrides an explicit text (@) or custom number format",
        isinstance(ws["A1"].value, (int, float))
        or (ws["A2"].number_format != '0.00" kg"'
            and '0.00" kg"' not in _f09_notes),
        f'A1 (was format "@") holds {ws["A1"].value!r} as '
        f'{type(ws["A1"].value).__name__} with format {ws["A1"].number_format!r}\n'
        f'A2 (was \'0.00" kg"\') now has format {ws["A2"].number_format!r} and no note '
        f'names the pattern it replaced')

# F-10 — a value that stays TEXT written into a DATE-formatted cell is invisible: it
# displays exactly like the real dates around it, the card shows the raw serial, and the
# receipt says it matched.
def _b(wb):
    ws = wb.active
    ws["A1"] = _dt.date(2026, 1, 10)
    ws["A2"] = _dt.date(2026, 1, 11)
    ws["A3"] = _dt.date(2026, 1, 12)
    ws["B1"] = "=A3-A1"
wb_new("f10.xlsx", _b)
out, a, _r = sa("f10.xlsx", [{"op": "set", "at": "A3", "values": [["2026-01-15"]]}])
ws = load("f10.xlsx")
pv = (out["preview"] or [{}])[0]
finding("F-10", LIE,
        "text written into a date-formatted cell is never flagged; the card shows the "
        "raw serial and the receipt says match",
        isinstance(ws["A3"].value, str)
        and not any("date" in n.lower() for n in out["notes"]),
        f'A3 is now {ws["A3"].value!r} (str) with format {ws["A3"].number_format!r} — it '
        f'displays as a date and breaks =A3-A1\n'
        f'card before_display={pv.get("before_display")!r} (a serial, not the date) '
        f'after_display={pv.get("after_display")!r} coerced={pv.get("coerced")}\n'
        f'receipt verify={json.dumps((a or {}).get("verify"))}')

# F-11 — op_read hands a date cell to the model as a bare serial number.
wb_new("f11.xlsx", lambda wb: [wb.active.__setitem__("A1", _dt.date(2026, 1, 15)),
                               wb.active.__setitem__("A2", _dt.time(9, 30))])
r, _e = oo.op_read(R, "f11.xlsx")
cells = {c["ref"]: c for c in r["cells"]}
finding("F-11", LIE,
        "op_read reports a date/time cell as a bare serial with no format and no note",
        cells.get("A1", {}).get("text") == "46037",
        f'A1 → {json.dumps(cells.get("A1"))}   A2 → {json.dumps(cells.get("A2"))}\n'
        f'notes mention dates? {any("date" in n.lower() for n in r["notes"])}')

# F-12 — an integer past 2**53 is silently rounded, and the card shows a THIRD value.
wb_new("f12.xlsx")
out, a, _r = sa("f12.xlsx", [{"op": "set", "at": "A1",
                              "values": [[12345678901234567, "12345678901234567"]]}])
# FIXED means the staging REFUSED it with a sentence naming the limit, rather than
# writing a third value nobody asked for.
finding("F-12", LIE,
        "integers past 2**53 lose precision silently; the card prints a value that "
        "matches neither the request nor the disk",
        out is not None or "exactly" not in str(_r or ""),
        f'staging said {_r!r} and the workbook now holds '
        f'{load("f12.xlsx")["A1"].value!r}' if out is not None
        else f'refusal: {_r}')

# F-33 — a save destroys every cached formula result in the workbook, unmentioned.
def _cached(path):
    """Forge a cached <v> onto the formula, the way Excel/ONLYOFFICE would."""
    tmp = path + ".z"
    zin, zout = zipfile.ZipFile(path), zipfile.ZipFile(tmp, "w")
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename == "xl/worksheets/sheet1.xml":
            data = data.decode().replace("<f>SUM(A1:A2)</f>",
                                         "<f>SUM(A1:A2)</f><v>30</v>").encode()
        zout.writestr(it, data)
    zin.close()
    zout.close()
    shutil.move(tmp, path)


p = wb_new("f33.xlsx", lambda wb: [wb.active.__setitem__("A1", 10),
                                   wb.active.__setitem__("A2", 20),
                                   wb.active.__setitem__("A3", "=SUM(A1:A2)")])
_cached(p)
before, _e = oo.op_read(R, "f33.xlsx")
out, a, _r = sa("f33.xlsx", [{"op": "set", "at": "B1", "values": [["poke"]]}])
after, _e = oo.op_read(R, "f33.xlsx")


def _f(rd, ref):
    return next((c for c in rd["cells"] if c["ref"] == ref), {})


finding("F-33", COS,
        "an apply wipes every cached formula result in the workbook and nothing says so",
        # FIXED means the wipe is still there (it is openpyxl's, and unavoidable) but the
        # apply result now SAYS SO, which is the whole of the finding's "correct".
        ("cached_value" in _f(before, "A3") and "cached_value" not in _f(after, "A3")
         and not any("cach" in n.lower() for n in (a or {}).get("notes", []))),
        f'before apply: {json.dumps(_f(before, "A3"))}\n'
        f'after  apply: {json.dumps(_f(after, "A3"))}\n'
        f'apply notes mention it? '
        f'{any("cach" in n.lower() for n in (a or {}).get("notes", []))}')

# F-14 — "cached": true is emitted for a formula cell that has NO cached value.
finding("F-14", LIE,
        '"cached": true is emitted for formula cells that hold no cached value at all',
        _f(after, "A3").get("cached") is True and "cached_value" not in _f(after, "A3"),
        f'{json.dumps(_f(after, "A3"))} — the flag asserts a cached value the cell does '
        "not have")

# F-28 — FORCE_STRING (t=4) intent does not survive the round-trip, so the new
# text-that-looks-numeric warning will nag about cells that are text ON PURPOSE.
snap = office.empty_snapshot("f28")
snap["sheets"]["sheet-01"]["cellData"] = {"0": {"0": {"v": 7, "t": office.CV_FORCE_STRING}}}
office.write_snapshot(snap, os.path.join(OD, "f28.xlsx"))
snap2 = office.snapshot_from_path(os.path.join(OD, "f28.xlsx"))
cell = oo.cell_at(snap2["sheets"]["sheet-01"], 0, 0)
finding("F-28", COS,
        "FORCE_STRING intent is lost on the round-trip, so a deliberately-text cell "
        "becomes a text-that-looks-numeric warning",
        cell.get("t") == office.CV_STRING and oo.text_numeric(cell) is not None,
        f'wrote {{v:7,t:4}} → re-read {json.dumps(cell)} → '
        f'text_numeric says {json.dumps(oo.text_numeric(cell))}')


# ══ 2. FORMULAS ═══════════════════════════════════════════════════════════════
print("\n── 2. formulas ──")

# F-04 — the formula honesty note is scoped to the TARGET sheet. A row insert in Data
# silently breaks =SUM(Data!A1:A5) on Summary and nothing anywhere says so.
def _b(wb):
    d = wb.active
    d.title = "Data"
    for i in range(1, 6):
        d.cell(i, 1, i * 10)
    s = wb.create_sheet("Summary")
    s["A1"] = "=SUM(Data!A1:A5)"
    s["A2"] = "=Data!A3"
wb_new("f04.xlsx", _b)
out, a, _r = sa("f04.xlsx", [{"op": "insert", "what": "row", "at": "2"}], sheet="Data")
note = any("references were NOT rewritten" in n for n in out["notes"])
finding("F-04", LIE,
        "the formula honesty note is per-sheet: a cross-sheet reference breaks with no "
        "note on the card, the receipt or the tool result",
        not note,
        "Summary!A1 is still =SUM(Data!A1:A5) and now sums a shifted range; "
        "Summary!A2 (=Data!A3) points at different data.\n"
        f'card notes: {[n[:70] for n in out["notes"]]}')

# F-08 / Z1 — a formula with a trailing space makes the RECEIPT report a failed write.
wb_new("f08.xlsx")
out, a, _r = sa("f08.xlsx", [{"op": "set", "at": "A1", "values": [["=SUM(B1:B2) "]]}])
ws = load("f08.xlsx")
v = (a or {}).get("verify") or [{}]
finding("F-08", LIE,
        "a trailing space on a formula makes the receipt say the write failed when it "
        "was perfect",
        ws["A1"].value == "=SUM(B1:B2)" and v[0].get("match") is False,
        f'disk holds {ws["A1"].value!r} — correct — but verify says '
        f'{json.dumps(v[0])}\nverify_note: {(a or {}).get("verify_note")}\n'
        "the panel's success badge renders from these numbers")

# F-25 — the honesty note fires even when the op could not have touched any reference.
wb_new("f25.xlsx", lambda wb: [wb.active.__setitem__("A1", 1),
                               wb.active.__setitem__("B1", "=A1*2")])
out, _a, _r = sa("f25.xlsx", [{"op": "insert", "what": "row", "at": "900"}])
finding("F-25", COS,
        "the formula honesty note fires for an insert far below every formula and every "
        "reference",
        any("references were NOT rewritten" in n for n in out["notes"]),
        "inserted a row at 900 on a sheet whose only formula is B1=A1*2 — nothing could "
        "have moved, and the card still warns her to check her formulas")

# F-29 / F-30 — syntactically invalid and circular formula text is written verbatim.
wb_new("f29.xlsx")
out, a, _r = sa("f29.xlsx", [{"op": "set", "at": "A1",
                              "values": [["=="], ["=A1+"], ["=A3+1"]]}])
ws = load("f29.xlsx")
finding("F-29", COS,
        "syntactically invalid formula text is written verbatim as a formula with no "
        "note (Excel may report the workbook as needing repair)",
        ws["A2"].value == "=A1+"
        and not any("repair" in n.lower() for n in out["notes"]),
        f'A1={ws["A1"].value!r} A2={ws["A2"].value!r} '
        f'notes={[n[:60] for n in out["notes"]]}')
finding("F-30", COS,
        "a self-referencing formula is accepted with no note",
        ws["A3"].value == "=A3+1"
        and not any("circular" in n.lower() for n in out["notes"]),
        f'A3={ws["A3"].value!r} — Excel will pop a circular-reference dialog on open')


# ══ 3. STRUCTURE OPS ══════════════════════════════════════════════════════════
print("\n── 3. structure ops ──")

# F-03 / Z4 — rc_apply remaps only the first TIER1_MAX_COLS columns. On a wider sheet a
# row insert leaves columns past GR unshifted, a column insert DESTROYS a value, and the
# note that fires describes something that did not happen.
def _wide(wb):
    ws = wb.active
    for c in range(1, 251):
        ws.cell(1, c, f"h{c}")
        ws.cell(2, c, c)
wb_new("f03.xlsx", _wide)
out, _a, _r = sa("f03.xlsx", [{"op": "insert", "what": "row", "at": "1"}])
ws = load("f03.xlsx")
row_bad = ws.cell(1, 201).value == "h201" and ws.cell(1, 200).value is None
wb_new("f03b.xlsx", lambda wb: [wb.active.cell(1, c, c) for c in range(1, 251)])
out2, _a, _r = sa("f03b.xlsx", [{"op": "insert", "what": "col", "at": "A"}])
ws2 = load("f03b.xlsx")
col_bad = ws2.cell(1, 201).value == 201        # should be 200 after a left insert
finding("F-03", LIE,
        "insert/delete row or column silently mis-shifts (and can destroy) every column "
        "past GR on a sheet wider than 200 columns",
        row_bad or col_bad,
        f'ROW insert on a 250-col sheet: col GR row1={ws.cell(1, 200).value!r} (shifted) '
        f'but col GS row1={ws.cell(1, 201).value!r} (NOT shifted) — the header row is '
        "now split across two rows\n"
        f'COL insert on a 250-col sheet: col 200 was lost, col 201 still holds '
        f'{ws2.cell(1, 201).value!r}\n'
        f'the only note is: {[n for n in out["notes"] if "fell off" in n]}\n'
        "— which says one column fell off the edge, not that 50 stopped moving")

# F-03b — the same wrong note is used for a DELETE, and it says "the insert".
wb_new("f03c.xlsx", _wide)
out3, _a, _r = sa("f03c.xlsx", [{"op": "delete_rc", "what": "row", "at": "1"}])
finding("F-03b", COS,
        'the clipped-grid note says "the insert" on a delete',
        any(n.startswith("the insert reached") for n in out3["notes"]),
        f'{[n[:80] for n in out3["notes"] if "insert reached" in n]}')

# F-05 — merge changes are invisible in the preview.
def _b(wb):
    ws = wb.active
    ws["A1"] = "t"
    ws.merge_cells("A1:A4")
wb_new("f05.xlsx", _b)
out, a, _r = sa("f05.xlsx", [{"op": "delete_rc", "what": "row", "at": "2", "n": 2}])
ws = load("f05.xlsx")
finding("F-05", LIE,
        "the preview never mentions merged ranges, so a change that only reshapes a "
        "merge shows as 0 cells and no note",
        (not out.get("merge_changes")
         and "A1:A2" in [str(m) for m in ws.merged_cells.ranges]),
        f'card: {out["summary"]!r}  preview={out["preview"]} '
        f'merge_changes={out.get("merge_changes")}\n'
        f'the merge went A1:A4 → {[str(m) for m in ws.merged_cells.ranges]} and nothing '
        "on the card said a merge would change")

# F-18 — add_sheet infinite-loops when the 31-char name is already taken. THIS HANGS THE
# BRIDGE THREAD FOREVER. Guarded by SIGALRM so this file cannot hang.
NAME31 = "x" * office.SHEET_NAME_MAX
p = os.path.join(OD, "f18.xlsx")
wbk = openpyxl.Workbook()
wbk.create_sheet(NAME31)
wbk.save(p)
snap = office.snapshot_from_path(p)
hung = "skip"
if hasattr(signal, "SIGALRM"):
    def _boom(_s, _f):
        raise TimeoutError("hung")
    old = signal.signal(signal.SIGALRM, _boom)
    signal.setitimer(signal.ITIMER_REAL, 2.0)
    try:
        oo.add_sheet(snap, NAME31)
        hung = False
    except TimeoutError:
        hung = True
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, old)
finding("F-18", REF,
        "add_sheet / rename with a name already at the 31-char sheet limit loops "
        "FOREVER — the bridge thread never returns",
        hung,
        f'office_ops.add_sheet(snapshot, {NAME31!r}) never returns: the "Name 2" step is '
        "re-truncated to the same 31 characters, so the while-loop can never find a free "
        "name. Reached from BOTH stage_changes and apply_changeset.")

# F-19 — an Excel-invalid sheet name stages fine and 500s on Apply, after the checkpoint,
# the daily .bak and the pre-agent copy have all already been taken.
oo.changeset_clear()
wb_new("f19.xlsx", lambda wb: wb.active.__setitem__("A1", "important"))
out, a, reason = sa("f19.xlsx", [{"op": "set", "at": "A2", "values": [["also mine"]]},
                                 {"op": "add_sheet", "name": "Q1:Q2"}])
crashed = (a is None and str(reason).startswith("EXC")) or (
    out is not None and a is None)
finding("F-19", REF,
        "a sheet name with an Excel-invalid character ( : [ ] * ? / \\ ) stages happily "
        "and raises an uncaught ValueError on Apply — HTTP 500, no sentence",
        crashed,
        (f'stage: OK, card says {out["op_list"]}\napply: {reason}\n'
         f'and it left a checkpoint behind: {len(oo.list_checkpoints(R, "f19.xlsx"))}, '
         f'plus {sorted(x for x in os.listdir(OD) if x.startswith("f19"))}\n'
         "the whole changeset (including the set) is lost, and every retry crashes again")
        if out is not None else f'refused at staging: {reason}')

# F-15 — a rename whose `at` names no sheet renames the FIRST sheet, and the card never
# says which sheet it is renaming.
def _two(wb):
    wb.active.title = "Alpha"
    wb.create_sheet("Beta")
wb_new("f15.xlsx", _two)
out, a, _r = sa("f15.xlsx", [{"op": "sheet", "rename": "Renamed", "at": "Nonexistent"}])
names = openpyxl.load_workbook(os.path.join(OD, "f15.xlsx")).sheetnames
finding("F-15", LIE,
        "a rename whose `at` names no existing sheet silently renames the FIRST sheet, "
        "and op_summary never names the sheet at all",
        names[0] == "Renamed",
        f'sheets went ["Alpha","Beta"] → {names}\n'
        f'the card only said: {out["op_list"]}  preview={out["preview"]}  '
        f'notes={[n[:50] for n in out["notes"] if "sheet" in n.lower()]}')

# F-16 — a rename onto an existing sheet name silently renames the OTHER sheet.
wb_new("f16.xlsx", _two)
out, a, _r = sa("f16.xlsx", [{"op": "sheet", "rename": "Beta"}])
names = openpyxl.load_workbook(os.path.join(OD, "f16.xlsx")).sheetnames
finding("F-16", LIE,
        "renaming a sheet onto a name another sheet already has silently renames THAT "
        "sheet to Name(2)",
        "Beta(2)" in names,
        f'sheets went ["Alpha","Beta"] → {names} — Debi\'s own "Beta" sheet was renamed '
        "by a change that only offered to rename one sheet")

# F-31 — the write-side dedup can emit a 32-character sheet title.
snap = office.empty_snapshot("f31")
for i in range(2, 14):
    sid = f"s{i}"
    snap["sheets"][sid] = dict(snap["sheets"]["sheet-01"], id=sid, name="Z" * 31)
    snap["sheetOrder"].append(sid)
long_titles = [n for _sid, n in office._sheet_names(snap)
               if len(n) > office.SHEET_NAME_MAX]
finding("F-31", COS,
        "office._sheet_names' duplicate step can produce a sheet title longer than "
        "Excel's 31 characters",
        bool(long_titles),
        f'{len(long_titles)} title(s) over the limit, e.g. {long_titles[:1]} — openpyxl '
        "warns and some applications refuse the file")

# F-24 — a destructive `n` is coerced to 1 instead of refused.
def _n(v):
    ops, _c, _r = oo.validate_ops([{"op": "delete_rc", "what": "row", "at": "1", "n": v}])
    return (ops or [{}])[0].get("n")
bad = {v: _n(v) for v in (0, -5, 1.9, "all", None, "3")}
finding("F-24", COS,
        'a delete\'s row count silently becomes 1 for n=0, n=-5, n="all"',
        bad[0] == 1 and bad["all"] == 1,
        f'{bad} — a model that said n:"all" and reported "deleted the rows" is not '
        "contradicted by anything except the card's own op line")


# ══ 4. THE STAGING / APPLY LIFECYCLE ══════════════════════════════════════════
print("\n── 4. staging / apply lifecycle ──")

# F-01 — THE MTIME FENCE IS RECORDED AND NEVER CHECKED. stage → the file changes on disk
# (the editor saved, another apply landed, Debi edited it) → apply runs anyway, against a
# preview that is now fiction.
wb_new("f01.xlsx", lambda wb: [wb.active.cell(i, 1, i * 10) for i in range(1, 6)])
out, _reason = stage("f01.xlsx", [{"op": "sort", "col": "A", "order": "desc"}])
staged_mtime = oo._CHANGESETS[out["changeset_id"]]["file_mtime"]
w = openpyxl.load_workbook(os.path.join(OD, "f01.xlsx"))
for i in range(1, 6):
    w.active.cell(i, 1, 1000 - i)
w.active["B1"] = "her own note"
w.save(os.path.join(OD, "f01.xlsx"))
a, reason = apply(out["changeset_id"])
finding("F-01", LIE,
        "apply_changeset records the file's mtime at staging and NEVER CHECKS IT — a "
        "changeset applies over a workbook that changed since the preview",
        a is not None,
        f'staged against mtime {staged_mtime:.3f}; applied against '
        f'{os.path.getmtime(os.path.join(OD, "f01.xlsx")):.3f}\n'
        f'the card promised A1: 10 → 50. The receipt now reports '
        f'{(a or {}).get("verify_note")}\n'
        "office_ops.py's own UNDO path has exactly this fence (FENCE_EPS / "
        "UNDO_FENCE_REFUSAL) — the APPLY path does not")

# F-32 — apply after a rename or a delete refuses with a bare "no such workbook".
wb_new("f32.xlsx", lambda wb: wb.active.__setitem__("A1", 1))
out, _reason = stage("f32.xlsx", [{"op": "set", "at": "A1", "values": [[2]]}])
office.rename_doc(R, "f32.xlsx", "f32-renamed.xlsx")
_a, reason_ren = apply(out["changeset_id"])
wb_new("f32b.xlsx", lambda wb: wb.active.__setitem__("A1", 1))
out2, _reason = stage("f32b.xlsx", [{"op": "set", "at": "A1", "values": [[2]]}])
os.remove(os.path.join(OD, "f32b.xlsx"))
_a2, reason_del = apply(out2["changeset_id"])
finding("F-32", COS,
        'apply after a rename or a delete answers "no such workbook" — it never says '
        "the workbook moved or that the proposal is still pending",
        reason_ren == "no such workbook" and reason_del == "no such workbook",
        f'renamed → {reason_ren!r}\ndeleted → {reason_del!r}')

# F-21 — a rename orphans the checkpoint stack, so "Undo this change" dies for good.
wb_new("f21.xlsx", lambda wb: wb.active.__setitem__("A1", "orig"))
out, _reason = stage("f21.xlsx", [{"op": "set", "at": "A1", "values": [["new"]]}])
a, _r = apply(out["changeset_id"])
office.rename_doc(R, "f21.xlsx", "f21-renamed.xlsx")
u, reason = oo.undo_changeset(R, out["changeset_id"])
finding("F-21", REF,
        "renaming a workbook orphans its checkpoint stack and its pre-agent copy — undo "
        "is gone forever and the refusal blames a missing workbook",
        u is None,
        f'undo → {reason!r}\n'
        f'the checkpoint is still on disk under the OLD stem: '
        f'{sorted(os.listdir(os.path.join(OD, ".checkpoints")))}\n'
        "office.rename_doc moves neither .checkpoints/<stem>/ nor <stem>.pre-agent.xlsx")

# F-20 — a REFUSED apply still pushed a checkpoint, so mashing Apply while the editor is
# dirty prunes away every real undo point.
wb_new("f20.xlsx", lambda wb: wb.active.__setitem__("A1", "v0"))
real_ids = []
for i in range(3):
    o, _reason = stage("f20.xlsx", [{"op": "set", "at": "A1", "values": [[f"v{i+1}"]]}])
    aa, _r = apply(o["changeset_id"])
    real_ids.append(o["changeset_id"])
oo.heartbeat("f20.xlsx", True)                # Debi has the workbook open and dirty
before_n = len(oo.list_checkpoints(R, "f20.xlsx"))
for _i in range(10):
    o, _reason = stage("f20.xlsx", [{"op": "set", "at": "A1", "values": [["nope"]]}])
    _aa, refusal = apply(o["changeset_id"])
after_n = len(oo.list_checkpoints(R, "f20.xlsx"))
survivors = [c["changeset_id"] for c in oo.list_checkpoints(R, "f20.xlsx")]
oo.heartbeat_clear()
finding("F-20", REF,
        "an Apply refused for unsaved edits has ALREADY pushed a checkpoint — ten "
        "refused presses prune away every real undo point",
        not any(cid in survivors for cid in real_ids),
        f'3 real applies → {before_n} checkpoints; 10 REFUSED applies (each answering '
        f'{refusal[:40]!r}…) → {after_n} checkpoints, none of them from a real apply.\n'
        "_apply takes push_checkpoint's copy before the dirty check clears")

# F-07 — the pre-agent copy overwrites a real user file that happens to be called that.
wb_new("f07.pre-agent.xlsx", lambda wb: wb.active.__setitem__("A1", "MY REAL DATA"))
wb_new("f07.xlsx", lambda wb: wb.active.__setitem__("A1", "main"))
_out, a, _r = sa("f07.xlsx", [{"op": "set", "at": "A1", "values": [["changed"]]}])
victim = openpyxl.load_workbook(os.path.join(OD, "f07.pre-agent.xlsx")).active["A1"].value
finding("F-07", LIE,
        "the pre-agent copy silently destroys a real workbook that happens to be named "
        "<stem>.pre-agent.xlsx",
        victim != "MY REAL DATA",
        f'f07.pre-agent.xlsx held "MY REAL DATA" and now holds {victim!r}\n'
        f'op_list lists it as an ordinary workbook, so Debi can create exactly this '
        "name from the panel; nothing on the card or the receipt names the file it is "
        "about to overwrite")

# F-22 — an apply on a workbook literally named *.pre-agent.xlsx refuses with raw
# shutil text.
out, a, reason = sa("f07.pre-agent.xlsx",
                    [{"op": "set", "at": "A2", "values": [["x"]]}])
finding("F-22", REF,
        "a workbook named <stem>.pre-agent.xlsx can never be written: the refusal is a "
        "raw shutil SameFileError string",
        a is None and "same file" in str(reason).lower(),
        f'{reason}')

# F-17 — the changeset store cap silently evicts pending proposals, four at a time, and
# never tells the session whose proposal it threw away.
oo.changeset_clear()
ids = []
# ⚠️ THE WORKBOOKS ARE BUILT WITHOUT wb_new ON PURPOSE: wb_new calls changeset_clear(),
# which also empties the SESSION-LINE queue this probe is about, so using it inside the
# loop would erase the very evidence being checked.
for i in range(oo.CHANGESET_MAX + 6):
    nm = f"f17-{i}.xlsx"
    _w = openpyxl.Workbook()
    _w.active["A1"] = 1
    _w.save(os.path.join(OD, nm))
    o, _reason = oo.stage_changes(R, f"sess{i}", nm, None,
                                  [{"op": "set", "at": "A1", "values": [[2]]}])
    ids.append((o["changeset_id"], nm, f"sess{i}"))
gone = [(nm, s) for cid, nm, s in ids if not oo.get_changeset(cid)]
told = [s for _nm, s in gone if oo.peek_session_lines(s)]
finding("F-17", LIE,
        "the changeset store cap evicts pending proposals silently — the model was told "
        "the change was staged and the card is simply gone",
        bool(gone) and not told,
        f'staged {len(ids)}, {len(gone)} evicted, {len(told)} of those sessions were '
        "told anything\n"
        "stage_changes drops the 4 oldest whenever the store reaches CHANGESET_MAX; "
        "no session line, no note, nothing in the panel")


# ══ 5. ROUND-TRIP HONESTY ═════════════════════════════════════════════════════
print("\n── 5. round-trip honesty ──")

# F-27 — four things the fidelity contract does not name are lost anyway, and the
# snapshot emits FAKE values for three of them.
def _fid(wb):
    ws = wb.active
    ws.title = "Kept"
    ws["A1"] = "x"
    ws["E1"] = "y"
    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["D"].hidden = True
    ws.row_dimensions[11].hidden = True
    ws.auto_filter.ref = "A1:E10"
p = wb_new("f27.xlsx", _fid)


def _feats(path):
    ws = openpyxl.load_workbook(path)["Kept"]
    return {"freeze": ws.freeze_panes, "gridlines": ws.sheet_view.showGridLines,
            "col_D_hidden": bool(ws.column_dimensions["D"].hidden),
            "row_11_hidden": bool(ws.row_dimensions[11].hidden),
            "autofilter": ws.auto_filter.ref}


b4 = _feats(p)
sa("f27.xlsx", [{"op": "set", "at": "Z1", "values": [["poke"]]}], sheet="Kept")
af = _feats(p)
snap = office.snapshot_from_path(p)
sh = snap["sheets"]["sheet-01"]
finding("F-27", LIE,
        "freeze panes, hidden rows/columns, gridline visibility and autofilter are lost "
        "by every apply — none of them is in the fidelity contract, and the snapshot "
        "reports FAKE values for them",
        b4 != af,
        f'before {json.dumps(b4)}\nafter  {json.dumps(af)}\n'
        f'a hidden column becoming visible is data Debi hid coming back on screen.\n'
        f'the snapshot asserts freeze={json.dumps(sh.get("freeze"))} '
        f'showGridlines={sh.get("showGridlines")} and hd=0 on every row/column entry — '
        "read from nothing and written to nothing, so any reader that trusts them is "
        "wrong.\n"
        f'FIDELITY_NOTE names none of them: {office.FIDELITY_NOTE!r}')

# F-06 — a style-only changeset has an EMPTY preview, claims 0 cells change, and gets a
# green "0 of 0 re-read cells" receipt.
wb_new("f06.xlsx", lambda wb: [wb.active.cell(r, 1, r) for r in range(1, 6)])
out, a, _r = sa("f06.xlsx", [{"op": "style", "at": "A1:A5",
                              "set": {"bl": 1, "bg": "#ff0000"}}])
ws = load("f06.xlsx")
finding("F-06", LIE,
        "a formatting-only changeset previews as nothing: 0 cells changed, empty "
        "before→after, and a receipt that verified 0 of 0 cells",
        # FIXED means the card has formatting before → after rows AND the receipt
        # actually re-read them — both halves of the finding's "correct behaviour".
        (not out.get("style_preview")
         or not [v for v in ((a or {}).get("verify") or [])
                 if v.get("kind") == "format"]) and ws["A1"].font.bold,
        f'card: {out["summary"]!r} cells_changed={out["cells_changed"]} '
        f'preview={out["preview"]} style_preview={out.get("style_preview")}\n'
        f'receipt: {(a or {}).get("verify_note")!r} over '
        f'{len((a or {}).get("verify") or [])} rows\n'
        f'and the write DID happen: A1 bold={ws["A1"].font.bold} '
        f'fill={ws["A1"].fill.start_color.rgb}\n'
        "snapshot_diff compares cell_face, which ignores styles entirely")

# F-23 — a sheet-add or sheet-rename changeset verifies nothing and says so nonsensically.
wb_new("f23.xlsx", lambda wb: wb.create_sheet("Data"))
out, a, _r = sa("f23.xlsx", [{"op": "add_sheet", "name": "Data"}])
finding("F-23", REF,
        'a sheet-add changeset gets the receipt "0 of 0 re-read cell(s) hold what the '
        'change said they would" and verifies nothing at all',
        (a or {}).get("verify") == []
        and "0 of 0" in str((a or {}).get("verify_note")),
        f'sheets_added={(a or {}).get("sheets_added")} (the name collided and stepped, '
        f'which is right) but verify={json.dumps((a or {}).get("verify"))} '
        f'verify_note={(a or {}).get("verify_note")!r}\n'
        "nothing re-reads the workbook to confirm the sheet exists")


# ══ 6. CAPS AND REFUSALS ══════════════════════════════════════════════════════
print("\n── 6. caps and refusals (these all behave — recorded so a fix cannot regress "
      "them) ──")
CAPS = [
    ("60 ops", [{"op": "set", "at": "A1", "values": [[1]]}] * 60, True),
    ("61 ops", [{"op": "set", "at": "A1", "values": [[1]]}] * 61, False),
    ("2000 cells", [{"op": "set", "at": "A1",
                     "values": [[1] * 100 for _ in range(20)]}], True),
    ("2001 cells", [{"op": "set", "at": "A1",
                     "values": [[1] * 100 for _ in range(20)]},
                    {"op": "set", "at": "A30", "values": [[1]]}], False),
    ("row 5000", [{"op": "set", "at": "A5000", "values": [[1]]}], True),
    ("row 5001", [{"op": "set", "at": "A5001", "values": [[1]]}], False),
    ("col GR", [{"op": "set", "at": "GR1", "values": [[1]]}], True),
    ("col GS", [{"op": "set", "at": "GS1", "values": [[1]]}], False),
    ("style GR5000", [{"op": "style", "at": "GR5000", "set": {"bl": 1}}], True),
    ("style GS5000", [{"op": "style", "at": "GS5000", "set": {"bl": 1}}], False),
    ("delete 200 rows", [{"op": "delete_rc", "what": "row", "at": "1", "n": 200}], True),
    ("delete 201 rows", [{"op": "delete_rc", "what": "row", "at": "1", "n": 201}], False),
]
wrong = []
for label, ops, want_ok in CAPS:
    got, _c, reason = oo.validate_ops(ops)
    if bool(got) != want_ok:
        wrong.append(f"{label}: expected {'accept' if want_ok else 'refuse'}, "
                     f"got {reason or 'accept'}")
finding("CTRL-CAPS", CTL, "cap boundaries (60/2000/GR5000/RC_MAX, exactly at and one over)",
        bool(wrong), "\n".join(wrong))

REFS = [("A0", False), ("1A", False), ("A1:", False), (":B2", False), ("A", False),
        ("A1:B2:C3", False), ("AAAA1", False), ("A99999999", False),
        ("D5:A1", True), ("$A$1", True), ("ZZZ99999", True)]
wrong = [f"{ref}: act_range={oo.act_range(ref)}"
         for ref, ok in REFS if bool(oo.act_range(ref)) != ok]
finding("CTRL-REFS", CTL, "malformed A1 references are refused; reversed ranges normalise",
        bool(wrong), "\n".join(wrong))

# all-or-nothing on a mixed op list
wb_new("aon.xlsx")
out, reason = stage("aon.xlsx", [{"op": "set", "at": "A1", "values": [["good"]]},
                                 {"op": "set", "at": "A0", "values": [["bad"]]}])
finding("CTRL-AON", CTL,
        "one invalid op refuses the whole changeset and writes nothing (all-or-nothing "
        "holds)",
        out is not None or load("aon.xlsx")["A1"].value is not None,
        f'{reason}')


# ══ 7. GROUNDING TRUTHFULNESS (read side) ═════════════════════════════════════
print("\n── 7. grounding truthfulness ──")

# F-02 — the CARD's `sheet` field is the sheet the MODEL asked for, not the one that will
# be written, and _sheet_note never reaches staging.
def _three(wb):
    wb.active.title = "Alpha"
    wb.create_sheet("Beta")
    wb.create_sheet("Gamma")
wb_new("f02.xlsx", _three)
out, a, _r = sa("f02.xlsx", [{"op": "set", "at": "A1", "values": [["v"]]}],
                sheet="Q3 Data")
finding("F-02", LIE,
        "the changeset card names the sheet the MODEL asked for, not the sheet that will "
        "be written, and carries no note about the mismatch",
        out["sheet"] == "Q3 Data"
        and not any("no sheet called" in n for n in out["notes"]),
        f'card sheet={out["sheet"]!r} but every preview row says '
        f'{out["preview"][0]["sheet"]!r}\n'
        f'the honest sentence exists (office_ops._sheet_note) and reaches only the '
        f'APPLY result, after the fact: '
        f'{[n[:60] for n in (a or {}).get("notes", []) if "no sheet" in n]}')

# F-13 — sheet_stats counts a formula cell as TEXT *and* as a FORMULA, omits it from
# every aggregate, and says every number here is computed from the file.
def _fx(wb):
    ws = wb.active
    for i in range(1, 6):
        ws.cell(i, 1, i)
        ws.cell(i, 2, f"=A{i}*2")
wb_new("f13.xlsx", _fx)
st, _e = oo.op_sheet_stats(R, "f13.xlsx")
colB = st["columns"][1]
finding("F-13", LIE,
        "sheet_stats counts a formula cell as TEXT as well as a FORMULA and reports no "
        "aggregate for the column, under a note claiming completeness",
        colB["text"] == 5 and colB["formulas"] == 5 and "sum" not in colB,
        f'column B is 5 formula cells and reads back as {json.dumps(colB)}\n'
        '— 5 text + 5 formulas is 10 things in a 5-cell column, "numbers: 0" reads as '
        '"this column is not numeric", and there is no sum/min/max at all')

# targeting across three sheets DOES hold on the read side — recorded so a fix cannot
# regress it.
wb_new("f07r.xlsx", _three)
bad = []
for want, expect, note in (("Alpha", "Alpha", False), ("Beta", "Beta", False),
                           ("Gamma", "Gamma", False), ("beta ", "Beta", False),
                           ("Delta", "Alpha", True), (None, "Alpha", False)):
    r, _e = oo.op_read(R, "f07r.xlsx", want)
    warned = any("no sheet called" in n for n in r["notes"])
    if r["sheet"] != expect or warned != note:
        bad.append(f"{want!r} → {r['sheet']!r} warned={warned}")
finding("CTRL-SHEET", CTL,
        "op_read / op_sheet_stats sheet targeting across three sheets, including the "
        "unknown-name fallback and its warning",
        bool(bad), "\n".join(bad))

# F-26 — the merge note in op_read names a tool that no longer exists.
def _mg(wb):
    ws = wb.active
    ws["A1"] = "spans"
    ws.merge_cells("A1:C1")
wb_new("f26.xlsx", _mg)
r, _e = oo.op_read(R, "f26.xlsx")
finding("F-26", COS,
        "op_read's merge note tells the model office_sort refuses the sheet — there is "
        "no office_sort tool any more",
        any("office_sort" in n for n in r["notes"]),
        f'{[n for n in r["notes"] if "office_sort" in n]}\n'
        "the real behaviour is that a `sort` op inside a changeset is SKIPPED with a "
        "note — a different sentence and a different surface")


# ══ the tally ═════════════════════════════════════════════════════════════════
print("\n" + "=" * 100)
order = [LIE, REF, COS, CTL]
total_r = total_f = total_s = 0
for sev in order:
    v = SEV.get(sev)
    if not v:
        continue
    total_s += v["skipped"]
    if sev != CTL:
        total_r += v["reproduces"]
        total_f += v["fixed"]
    if sev == CTL:
        print(f"{sev:16}  {v['fixed']:2} hold · {v['reproduces']:2} BROKEN")
        continue
    print(f"{sev:16}  {v['reproduces']:2} reproduce · {v['fixed']:2} fixed"
          + (f" · {v['skipped']} skipped" if v["skipped"] else ""))
print(f"{'TOTAL':16}  {total_r:2} reproduce · {total_f:2} fixed"
      + (f" · {total_s} skipped" if total_s else "")
      + f"   ({len(SEEN)} checks)")
print("\nSee docs/research/2026-08-28-adversarial-findings-server.md for the correct "
      "behaviour of each finding.")
print("THIS FILE IS A CATALOGUE, NOT A GATE — exiting 0 on purpose.")
try:
    shutil.rmtree(R)
except OSError:
    pass
sys.exit(0)

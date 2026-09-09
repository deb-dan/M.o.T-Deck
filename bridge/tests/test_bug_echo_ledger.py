"""BUG-ECHO LEDGER — repros for every BUG-rated finding of the 2026-08-28 bug-echo
sweep (docs/research/2026-08-28-bug-echo-sweep.md).

⚠️ THIS FILE IS NOT A GATE AND MUST NEVER BE WIRED INTO ONE. It EXITS 0 ALWAYS.

WHAT IT IS: the campaign-ledger pattern of test_office_adversarial.py, applied to the
bug-echo methodology (github.com/Terryc21/bug-echo): after this project's own bugs were
FIXED and proven real (F-01…F-33 all read `fixed` in test_office_adversarial.py as of
v1.5.19; the 2026-08-27 false-Done incident was fixed in v1.5.9), the codebase was swept
for OTHER copies of the same mistakes. Every finding rated BUG in the sweep report has
an executable repro here, same BE-nn id, same order. A batch-fix builder runs this file,
fixes, and runs it again to watch the lines flip to `fixed`. When every line says
`fixed`, delete this file (or promote the probes into the real gate suites).

    data/bridge-venv/bin/python bridge/tests/test_bug_echo_ledger.py

Each check prints ONE line:

    ECHO BE-02 [F-01 class ]  REPRODUCES  one-line title
    ECHO BE-02 [F-01 class ]  fixed       one-line title

Lines marked [CONTROL] print `holds` — they pin the SIBLING behaviour that is already
correct (the thing each echo was measured against), so a batch fix cannot regress the
original fix while chasing its echo. A CONTROL that prints BROKEN is a new bug.

STATUS after the 2026-08-28 batch fix (this file is not the gate — the gate tests are):
  · BE-02 FIXED — office.save_doc grew the fence (expect_mtime / force / unfenced, the
    409-shaped refusal), /api/office/save forwards it and logs an explicit UNFENCED skip,
    office_ops._apply passes the mtime it read. Gate: test_office_lane.py (the fence and
    its F-20 ordering, the route's wiring) and test_oo_lane.py (writeback's own three
    states, W-01). REMAINING HALF, NOT OURS THIS ROUND: bridge/panel/office.html still
    sends no `expect_mtime`, so the tier-1 Save and the Quick lane's sort are written
    UNFENCED-and-logged rather than fenced. One line each in `save()` and `actWrite()`
    (`expect_mtime: extSeen`); that page is fenced to another builder this round.
  · BE-03 FIXED — the checkpoint namespace is keyed by the DOCUMENT
    (office.checkpoint_key: a spreadsheet keeps its bare stem, other types carry their
    extension, a stem that already ends in a type mark carries its own), with a one-time
    migration in office.migrate_checkpoint_ns and a hands-off rule for destructive ops
    (checkpoint_dir_to_move). Gate: test_office_journey.py (the journey) and
    test_office_mcp.py (the key, the inverse, the migration, the corner found live).
  · BE-01 FIXED (2026-08-28, the SSE/panel slice that owns bridge/panel/index.html) —
    the chat lane's ONE tool_output handler (shared by the Hermes, Odysseus-agent and
    direct lanes) grew the is_error branch: a failed or denied tool renders a red ✗ chip
    carrying the tool's own sentence, drawn on the message HOLDER so the end-of-turn
    renderChatBody() cannot delete it (re-asserted by chatToolErrsRedraw after that
    render), and the step is marked so chatSummary can never stamp ✓ over it. Gate:
    bridge/tests/test_chat_toolerr.js — the three chip functions EXECUTED against a DOM
    shim, plus the handler, the all-designs token proof, and the two controls (the
    bridge's is_error mapping and office.html's own chip).

EVERY LINE IN THIS FILE NOW READS `fixed`. Per the note above, the probes have all been
promoted into real gate suites (named per finding) and this file has done its job — it
is kept for one more wave as the campaign record, then deletable.
"""
from __future__ import annotations

import os
import re
import shutil
import sys
import tempfile
import time

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)

from bridge import office, office_ops as oo   # noqa: E402

try:
    import openpyxl
except Exception as _e:                                          # noqa: BLE001
    print(f"SKIPPED ENTIRELY: openpyxl is not importable ({_e}) — run this with "
          "data/bridge-venv/bin/python")
    sys.exit(0)


COUNTS = {"reproduces": 0, "fixed": 0, "holds": 0, "broken": 0, "skipped": 0}


def echo(bid, klass, title, reproduces, detail=""):
    """One ledger line. `reproduces` truthy = the echo is still there."""
    if reproduces == "skip":
        COUNTS["skipped"] += 1
        state = "SKIPPED   "
    elif klass == "CONTROL":
        if reproduces:
            COUNTS["broken"] += 1
            state = "BROKEN    "
        else:
            COUNTS["holds"] += 1
            state = "holds     "
    elif reproduces:
        COUNTS["reproduces"] += 1
        state = "REPRODUCES"
    else:
        COUNTS["fixed"] += 1
        state = "fixed     "
    print(f"ECHO {bid} [{klass:11}] {state}  {title}")
    if detail and reproduces and reproduces != "skip":
        for line in str(detail).splitlines():
            print(f"                                   · {line}")


R = tempfile.mkdtemp(prefix="bug-echo-")
OD = os.path.join(R, "data", "office")
os.makedirs(OD, exist_ok=True)


def wb_new(name, a1=10):
    oo.changeset_clear()
    oo.heartbeat_clear()
    p = os.path.join(OD, name)
    if os.path.exists(p):
        os.remove(p)
    wb = openpyxl.Workbook()
    wb.active["A1"] = a1
    wb.save(p)
    return p


def read_a1(name):
    return openpyxl.load_workbook(os.path.join(OD, name)).active["A1"].value


print(__doc__.splitlines()[0])
print("=" * 100)


# ══ BE-01 — the false-Done class (2026-08-27 consent incident) in the MAIN CHAT lane ══
# Original, proven and fixed: a tool failed, the panel rendered nothing, the model's
# "Done" was the only thing on screen. Fixed in v1.5.9 for the LOffice AI panel:
# routers/hermes.py maps `is_error` onto EVERY tool_output frame, and office.html draws
# a red ✗ chip with the tool's own sentence. The ECHO: index.html's Hermes chat lane
# consumes the same frames and renders `tool_output` as "reading results…"
# unconditionally — a failed tool is silent there, and a false "Done" stands exactly as
# it did in the incident. Source-level probe (the defect is a missing renderer branch).
print("\n── BE-01: false-Done class — the main-chat Hermes lane ──")
try:
    idx = open(os.path.join(ROOT, "bridge", "panel", "index.html"),
               encoding="utf-8", errors="replace").read()
    off = open(os.path.join(ROOT, "bridge", "panel", "office.html"),
               encoding="utf-8", errors="replace").read()
    hr = open(os.path.join(ROOT, "bridge", "routers", "hermes.py"),
              encoding="utf-8", errors="replace").read()

    # The bridge really does put is_error on the frame (so the fix is renderer-side).
    bridge_maps = 'fr["is_error"] = True' in hr
    # office.html really does render it (the original fix — pinned as the control).
    office_renders = ("is_error" in off) and ("csErrChip" in off)

    # index.html: find the SSE dispatch that handles 'tool_output' and see whether ANY
    # is_error read exists anywhere in the page. (As of the sweep: zero occurrences.)
    idx_handles_frames = "'tool_output'" in idx
    idx_reads_is_error = "is_error" in idx

    echo("BE-01", "false-Done",
         "index.html chat lane renders tool_output with no is_error branch — a failed "
         "tool is invisible and the model's narration stands uncontradicted",
         bridge_maps and idx_handles_frames and not idx_reads_is_error,
         "routers/hermes.py sets fr['is_error'] for every lane; office.html draws the\n"
         "✗ chip; index.html's handler (near 'tool_output' ~line 10527) says only\n"
         "'reading results…' whatever the frame carried.")
    # FIXED 2026-08-28. The probe above is deliberately left as-is (it is what proved
    # the bug), and these three lines are what the FIX has to keep true — a chip that
    # is drawn but deleted by the end-of-turn re-render would still pass a bare
    # `is_error in idx`, which is precisely the mistake worth pinning.
    echo("BE-01f", "false-Done",
         "…and the fix is a real chip, not just a mention of is_error: it is drawn on "
         "the holder, re-asserted after renderChatBody, and the step is marked so the "
         "✓ summary cannot contradict it",
         not ("function chatToolErrChip(" in idx
              and "chatToolErrsRedraw(holder);   // BE-01" in idx
              and "st.error" in idx))
    echo("BE-01c", "CONTROL",
         "the original v1.5.9 fix holds: the bridge maps is_error and office.html "
         "renders the ✗ chip",
         not (bridge_maps and office_renders))
except OSError as e:
    echo("BE-01", "false-Done", "source probe could not read the pages", "skip",
         str(e))


# ══ BE-02 — the F-01 class (fence recorded/available but not applied) at save_doc ═════
# Original, proven and fixed: apply_changeset staged an mtime and never checked it;
# the fix (v1.5.12) refuses when the workbook moved under the preview. Its siblings
# writeback (oo.py) and undo both fence. The ECHO: office.save_doc — the writer behind
# POST /api/office/save (the tier-1 grid's Save AND the Quick lane's sort route) — has
# no fence AT ALL: no parameter to express "the file I read at open time". A write that
# lands between the page's snapshot read and its save is silently clobbered, with only
# the once-per-day .bak behind it.
print("\n── BE-02: F-01 class — /api/office/save has no mtime fence ──")
wb_new("Race.xlsx", a1=10)
snap = office.snapshot_from_path(os.path.join(OD, "Race.xlsx"))   # page opens at T0
time.sleep(0.01)
w2 = openpyxl.load_workbook(os.path.join(OD, "Race.xlsx"))
w2.active["A1"] = 999                                             # someone else at T1
w2.save(os.path.join(OD, "Race.xlsx"))
rep, reason = office.save_doc(R, "Race.xlsx", snap)               # stale save at T2
clobbered = (reason is None and read_a1("Race.xlsx") != 999)
import inspect                                                    # noqa: E402
no_fence_param = "mtime" not in str(inspect.signature(office.save_doc))
echo("BE-02", "F-01 class",
     "save_doc overwrites a concurrent write silently — no expect_mtime parameter, "
     "no refusal, the sibling writers (writeback, apply, undo) all fence",
     clobbered and no_fence_param,
     f"A1 was 999 on disk when the stale snapshot saved; A1 now reads "
     f"{read_a1('Race.xlsx')!r}. save_doc signature: "
     f"{inspect.signature(office.save_doc)}")

# CONTROL: the sibling fence that BE-02 is measured against still refuses.
try:
    from bridge import oo as oomod
    p = os.path.join(OD, "Race.xlsx")
    stale = os.path.getmtime(p) - 5
    body = open(p, "rb").read()
    r2, err2 = oomod.writeback(office, R, "Race.xlsx", body, stale)
    echo("BE-02c", "CONTROL",
         "oo.writeback still refuses a stale-mtime save with a 409",
         not (r2 is None and err2 and err2[0] == 409))
except Exception as e:                                            # noqa: BLE001
    echo("BE-02c", "CONTROL", "writeback control errored", True, str(e))


# ══ BE-03 — the F-21 class comes back through the stem: the checkpoint namespace is ══
#            extension-blind, so a rename of Budget.docx orphans Budget.xlsx's undo
# Original, proven and fixed: a rename orphaned `.checkpoints/<stem>/` (F-21); the fix
# moves the folder with the file. The ECHO: `checkpoint_dir`, `pre_agent_for` and the
# rename's courtesy move all key by STEM ONLY, while three document types share the
# folder. Renaming a Word file that happens to share a stem with a spreadsheet moves
# the SPREADSHEET's checkpoint stack and pre-agent copy to the Word file's new stem —
# and the spreadsheet's undo then refuses with a sentence that blames pruning.
print("\n── BE-03: F-21 class — the undo stack is keyed by stem, blind to type ──")
wb_new("Budget.xlsx", a1=1)
out, _r = oo.stage_changes(R, "s", "Budget.xlsx", None,
                           [{"op": "set", "at": "A1", "values": [[2]]}])
rec = None
if out:
    rec, _r2 = oo.apply_changeset(R, out["changeset_id"])
ckdir = os.path.join(OD, office.CHECKPOINT_DIR, "Budget")
had_stack = bool(rec) and os.path.isdir(ckdir) and os.listdir(ckdir)
if had_stack:
    # A Word document sharing the stem. rename_doc never inspects bytes, so any bytes
    # under the right name exercise the rename path.
    shutil.copy2(os.path.join(OD, "Budget.xlsx"), os.path.join(OD, "Budget.docx"))
    new, r3 = office.rename_doc(R, "Budget.docx", "Notes.docx")
    moved = (new == "Notes.docx" and not os.path.isdir(ckdir)
             and os.path.isdir(os.path.join(OD, office.CHECKPOINT_DIR, "Notes")))
    u, r4 = oo.undo_changeset(R, out["changeset_id"])
    echo("BE-03", "F-21 class",
         "renaming a same-stem .docx moves the .xlsx's checkpoint stack with it; the "
         "spreadsheet's Undo then refuses, blaming the 10-checkpoint prune",
         moved and u is None,
         f".checkpoints/Budget → .checkpoints/Notes on a DOCX rename; "
         f"undo answered: {r4!r}")
else:
    echo("BE-03", "F-21 class", "could not build the checkpoint fixture", "skip",
         f"stage/apply failed: {_r!r}")

# CONTROL: the F-21 fix itself still holds for a SAME-TYPE rename.
wb_new("Ledger.xlsx", a1=1)
out2, _ = oo.stage_changes(R, "s", "Ledger.xlsx", None,
                           [{"op": "set", "at": "A1", "values": [[2]]}])
ok_ctrl = False
if out2:
    rec2, _ = oo.apply_changeset(R, out2["changeset_id"])
    if rec2:
        new2, _ = office.rename_doc(R, "Ledger.xlsx", "Ledger2.xlsx")
        u2, _ = oo.undo_changeset(R, out2["changeset_id"])
        ok_ctrl = (new2 == "Ledger2.xlsx" and u2 is not None)
echo("BE-03c", "CONTROL",
     "the F-21 fix holds: a same-type rename carries the undo stack with the file",
     not ok_ctrl)


# ══ summary ═══════════════════════════════════════════════════════════════════
print("=" * 100)
print(f"BUG echoes        {COUNTS['reproduces']} reproduce · {COUNTS['fixed']} fixed"
      f" · {COUNTS['skipped']} skipped")
print(f"CONTROLS          {COUNTS['holds']} hold · {COUNTS['broken']} BROKEN")
print("\nSee docs/research/2026-08-28-bug-echo-sweep.md for the full sweep "
      "(WATCH/REVIEW/OK findings live there — only BUG ratings have repros here).")
print("THIS FILE IS A LEDGER, NOT A GATE — exiting 0 on purpose.")
shutil.rmtree(R, ignore_errors=True)
sys.exit(0)
